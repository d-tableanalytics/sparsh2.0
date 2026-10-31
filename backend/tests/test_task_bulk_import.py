"""Delegation bulk import -- verification harness.

The spreadsheet half of the feature: the template (dropdowns of the user's categories and
people, mandatory columns marked), reading a filled file back, values filled in on the preview,
and turning each row into the Assign Task form's payload or a list of row errors. (Creating
the tasks is calendar_events.create_event, unchanged, and is not re-tested here.)

House convention: self-contained, no pytest, no database, ASCII output, exit 1 on fail.
Run:  python tests/test_task_bulk_import.py   (from backend/)
"""
import io
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

results = []


def check(label, condition, detail=None):
    results.append(bool(condition))
    print(f"  {'PASS' if condition else 'FAIL'}  {label}" + ("" if condition else f"  -- {detail}"))


def section(title):
    print(f"\n-- {title} --")


def main() -> int:
    from fastapi import HTTPException
    from openpyxl import load_workbook

    from app.services import task_bulk_import_service as B

    TODAY = date(2026, 10, 31)
    users = [{"_id": "u-asha", "full_name": "Asha Rao", "email": "asha@co.com"},
             {"_id": "u-ravi", "full_name": "Ravi", "email": "ravi@co.com"},
             {"_id": "u-mush", "full_name": "Mushkan", "email": "mushkan@co.com"},
             {"_id": "u-am1", "full_name": "Amit", "email": "amit1@co.com"},
             {"_id": "u-am2", "full_name": "Amit", "email": "amit2@co.com"}]
    assignable = B.People(users)
    everyone = B.People(users + [{"_id": "u-md", "full_name": "MD", "email": "md@co.com"}])

    section("1. Template built for the user")
    raw = B.build_template(["Reporting", "Marketing"], users)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Tasks"]
    headers = [c.value for c in ws[1]]
    check("sheets: Tasks, Instructions, Lists", wb.sheetnames == ["Tasks", "Instructions", "Lists"],
          wb.sheetnames)
    check("mandatory columns marked * : Task Title, Category, Deadline Date",
          [h for h in headers if h.endswith(" *")] == ["Task Title *", "Category *", "Deadline Date *"],
          headers)
    red = {c.value for c in ws[1] if c.fill.start_color.rgb.endswith("C00000")}
    check("mandatory headers are red", red == {"Task Title *", "Category *", "Deadline Date *"}, red)
    lists = wb["Lists"]
    check("Lists sheet holds the categories and 'Name <email>' people",
          [c.value for c in lists["A"] if c.value][1:] == ["Marketing", "Reporting"]
          and "Mushkan <mushkan@co.com>" in [c.value for c in lists["B"]], [c.value for c in lists["B"]])
    dv = {str(d.sqref).split(":")[0]: d for d in ws.data_validations.dataValidation}
    cat_col = "C"
    check("Category is a dropdown of the Lists sheet that still accepts a typed value",
          dv[f"{cat_col}2"].formula1.startswith("=Lists!$A$2") and dv[f"{cat_col}2"].errorStyle == "information")
    check("Assign To is a dropdown of the people", dv["D2"].formula1.startswith("=Lists!$B$2"))
    check("Deadline Date is a text cell with a DD-MM-YYYY hint and no Excel date rule "
          "(a month-day-year Excel must not reject 22-11-2026)",
          dv["F2"].type is None and not dv["F2"].showErrorMessage and "DD-MM-YYYY" in dv["F2"].prompt
          and ws["F2"].number_format == "@", (dv["F2"].type, ws["F2"].number_format))

    section("2. Reading a filled file")
    keys = [k for k, *_r in B.COLUMNS]
    def put(r, **values):
        for k, v in values.items():
            ws.cell(r, keys.index(k) + 1, v)
    put(2, title="Pay vendors", description="Q4 batch", category="Finance",
        assignees="Asha Rao <asha@co.com>, ravi@co.com", assignment_mode="Separate",
        deadline_date="05-11-2026", deadline_time="17:00", priority="High", watchers="md@co.com",
        tags="q4, vendors", checklist="Collect; Verify; Pay", evidence_required="Yes",
        verification_required="No")
    put(4, title="Self note", category="Admin", deadline_date="2026-11-02")
    buf = io.BytesIO(); wb.save(buf)
    rows = B.read_sheet("tasks.xlsx", buf.getvalue())
    check("two task rows read (blank row and pre-formatted empty rows skipped), row numbers kept",
          [r["_row"] for r in rows] == [2, 4], [r.get("_row") for r in rows])
    old = B.read_sheet("t.csv", b"Task Title,Category,Assign To (emails),Deadline Date\nCall,Sales,a@b.c,2026-11-02\n")
    check("a sheet with the old headers still reads", old[0]["assignees"] == "a@b.c")
    for name, data, frag in (("t.pdf", b"x", "xlsx"), ("t.csv", b"Name,Due\nx,y\n", "Task Title"),
                             ("t.csv", b"", "empty")):
        try:
            B.read_sheet(name, data); check(f"refuses {name} ({frag})", False)
        except HTTPException as e:
            check(f"refuses {name}: {e.detail[:60]}", frag.lower() in e.detail.lower(), e.detail)

    section("3. A valid row becomes the form's payload")
    res = B.validate_rows(rows, assignable, everyone, today=TODAY)
    p = res[0]["payload"]
    check("no errors on the full row", res[0]["errors"] == [], res[0]["errors"])
    check("'Name <email>' and a bare email both resolve; mode Separate",
          p["target_staff_id"] == ["u-asha", "u-ravi"] and p["assignment_mode"] == "separate", p)
    check("deadline 05-11-2026 17:00 IST -> 11:30 UTC", p["end"].startswith("2026-11-05T11:30"), p["end"])
    check("priority, tags, watchers, flags, checklist carried over",
          p["priority"] == "High" and p["tags"] == ["q4", "vendors"] and p["watchers"] == ["u-md"]
          and p["evidence_required"] is True and [c["title"] for c in p["checklist"]] == ["Collect", "Verify", "Pay"], p)
    q = res[1]["payload"]
    check("no assignee = myself; deadline still set", q and q["assigned_to"] == "myself"
          and q["end"].startswith("2026-11-02"), res[1])

    section("4. Names, and the mandatory deadline")
    r = B.validate_rows([{"_row": 2, "title": "task 31 oct", "category": "marketing",
                          "assignees": "Mushkan", "deadline_date": "03-11-2026"}],
                        assignable, everyone, today=TODAY)[0]
    rc = B.validate_rows([{"_row": 2, "title": "T", "category": "marketing", "deadline_date": "03-11-2026"}],
                         assignable, everyone, today=TODAY, categories=["Marketing", "HR"])[0]
    check("a typed 'marketing' uses the existing 'Marketing' category",
          rc["payload"]["category"] == "Marketing", rc["category"])
    check("a plain name ('Mushkan') resolves to the one person with that name",
          r["errors"] == [] and r["payload"]["target_staff_id"] == ["u-mush"], r)
    r = B.validate_rows([{"_row": 2, "title": "T", "category": "C", "assignees": "Amit",
                          "deadline_date": "03-11-2026"}], assignable, everyone, today=TODAY)[0]
    check("a name two people share asks for the email", "use the email" in " ".join(r["errors"]), r["errors"])
    r = B.validate_rows([{"_row": 2, "title": "task 31 oct", "category": "marketing", "assignees": "Mushkan"}],
                        assignable, everyone, today=TODAY)[0]
    check("no deadline -> 'Deadline Date is required', even for a valid assignee",
          len(r["errors"]) == 1 and r["errors"][0].startswith("Deadline Date is required"), r["errors"])

    section("5. Filling in a missing deadline on the preview")
    rows = [{"_row": 2, "title": "task 31 oct", "category": "marketing", "assignees": "Mushkan"},
            {"_row": 3, "title": "other", "category": "", "deadline_date": "03-11-2026"}]
    B.apply_overrides(rows, {"2": {"deadline_date": "2026-11-04", "deadline_time": "18:00",
                                   "title": "renamed"},
                             3: {"category": "Reporting", "assignees": ""}})
    fixed = B.validate_rows(rows, assignable, everyone, today=TODAY)
    check("the preview's deadline is used", fixed[0]["errors"] == []
          and fixed[0]["payload"]["end"].startswith("2026-11-04T12:30"), fixed[0])
    check("only editable columns are overridden (title kept)", fixed[0]["title"] == "task 31 oct")
    check("a category filled on the preview fixes that row", fixed[1]["errors"] == [], fixed[1]["errors"])

    section("6. Other row errors")
    bad = B.validate_rows([
        {"_row": 2, "title": "", "category": "", "deadline_date": "03-11-2026"},
        {"_row": 3, "title": "T", "category": "C", "assignees": "nobody@co.com", "deadline_date": "03-11-2026"},
        {"_row": 4, "title": "T", "category": "C", "deadline_date": "30-10-2026"},
        {"_row": 5, "title": "T", "category": "C", "deadline_date": "31-02-2026", "deadline_time": "25:00"},
        {"_row": 6, "title": "T", "category": "C", "deadline_date": "03-11-2026", "priority": "Urgent",
         "assignment_mode": "Both", "evidence_required": "maybe"},
        {"_row": 7, "title": "T", "category": "C", "deadline_date": "03-11-2026", "watchers": "ghost@co.com"},
    ], assignable, everyone, today=TODAY)
    err = {r["row"]: " | ".join(r["errors"]) for r in bad}
    check("missing title and category", "Title is required" in err[2] and "Category is required" in err[2], err[2])
    check("unknown assignee, shown as typed in the preview",
          "nobody@co.com (not found)" in err[3] and bad[1]["assignees"] == ["nobody@co.com"], err[3])
    check("past deadline refused", "cannot be in the past" in err[4], err[4])
    check("bad date and bad time", "not a date" in err[5] and "not a time" in err[5], err[5])
    check("bad priority, mode and Yes/No",
          "Priority" in err[6] and "Assignment Mode" in err[6] and "Evidence Required" in err[6], err[6])
    check("unknown In Loop", "In Loop not found" in err[7], err[7])
    check("a row with any error carries no payload (never half-created)",
          all(r["payload"] is None for r in bad if r["errors"]))

    section("6b. Checklist items")
    def ck(value):
        row = {"_row": 2, "title": "T", "category": "C", "deadline_date": "03-11-2026", "checklist": value}
        return [c["title"] for c in B.validate_rows([row], assignable, everyone, today=TODAY)[0]["payload"]["checklist"]]
    check("commas separate items: 'l1 ,l2' -> l1, l2", ck("l1 ,l2") == ["l1", "l2"], ck("l1 ,l2"))
    check("semicolons separate items", ck("A; B; C") == ["A", "B", "C"])
    check("new lines (Alt+Enter) separate items", ck("A\nB") == ["A", "B"])
    check("with ; present, a comma stays inside the item",
          ck("Collect data, verify; Share") == ["Collect data, verify", "Share"])

    section("7. Repeating tasks")
    rec = B.validate_rows([
        {"_row": 2, "title": "Weekly review", "category": "C", "assignees": "Asha Rao",
         "repeat": "Weekly", "repeat_on": "Mon, Thu", "start_date": "01-11-2026", "repeat_end_date": "31-03-2027"},
        {"_row": 3, "title": "Month close", "category": "C", "repeat": "monthly", "repeat_on": "1, 15"},
        {"_row": 4, "title": "Month end", "category": "C", "repeat": "Monthly", "repeat_on": "Last",
         "start_date": "05-11-2026"},
        {"_row": 5, "title": "Standup", "category": "C", "repeat": "Daily", "deadline_date": "10-11-2026"},
        {"_row": 6, "title": "Bad", "category": "C", "repeat": "Fortnightly"},
        {"_row": 7, "title": "Bad", "category": "C", "repeat": "Weekly", "repeat_on": "Funday"},
        {"_row": 8, "title": "Bad", "category": "C", "repeat": "Monthly", "repeat_on": "32"},
        {"_row": 9, "title": "Bad", "category": "C", "repeat": "Daily", "start_date": "01-12-2026",
         "repeat_end_date": "20-11-2026"},
        {"_row": 10, "title": "Bad", "category": "C", "repeat": "Daily", "start_date": "29-10-2026"},
        {"_row": 11, "title": "Yearly audit", "category": "C", "repeat": "Annually", "start_date": "15-11-2026"},
    ], assignable, everyone, today=TODAY)
    rr = {r["row"]: r for r in rec}
    w = rr[2]["payload"]
    check("Weekly Mon, Thu: weekdays sent in the form's numbering (Mon=1, Thu=4), no deadline",
          rr[2]["errors"] == [] and w["repeat"] == "Weekly" and w["repeat_data"]["weekdays"] == [1, 4]
          and w["end"] is None, rr[2])
    check("a Sunday start opens the series on the first Monday (02 Nov)",
          w["start"].startswith("2026-11-02") and w["repeat_end_date"] == "2027-03-31", w["start"])
    check("summary reads 'Weekly on Mon, Thu from 02 Nov 2026 to 31 Mar 2027'",
          rr[2]["repeat_summary"] == "Weekly on Mon, Thu from 02 Nov 2026 to 31 Mar 2027", rr[2]["repeat_summary"])
    m = rr[3]["payload"]
    check("Monthly 1, 15 with no Start Date starts today and opens on the 1st (01 Nov)",
          rr[3]["errors"] == [] and m["repeat_data"]["monthlyDates"] == [1, 15]
          and m["start"].startswith("2026-11-01") and m["repeat_end_date"] is None, rr[3])
    check("Monthly Last opens on the last day of the month (30 Nov)",
          rr[4]["payload"]["repeat_data"]["lastDay"] is True
          and rr[4]["payload"]["start"].startswith("2026-11-30"), rr[4].get("payload"))
    check("Daily ignores the Deadline Date column", rr[5]["errors"] == []
          and rr[5]["payload"]["end"] is None and rr[5]["payload"]["repeat"] == "Daily", rr[5])
    check("'Annually' is read as Yearly", rr[11]["payload"]["repeat"] == "Yearly", rr[11])
    check("unknown frequency", "Repeat must be" in " ".join(rr[6]["errors"]), rr[6]["errors"])
    check("unknown weekday", "unknown weekday" in " ".join(rr[7]["errors"]), rr[7]["errors"])
    check("month date out of range", "unknown month date" in " ".join(rr[8]["errors"]), rr[8]["errors"])
    check("end before first occurrence", "before the first occurrence" in " ".join(rr[9]["errors"]), rr[9]["errors"])
    check("past start date", "Start Date cannot be in the past" in " ".join(rr[10]["errors"]), rr[10]["errors"])

    passed = sum(results)
    print(f"\n=== {passed}/{len(results)} checks passed ===")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
