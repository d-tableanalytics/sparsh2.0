import os
import re
import pypdf
import docx
import json
import pandas as pd
import base64
from datetime import datetime
from bson import ObjectId
from app.db.mongodb import get_collection
from app.config.settings import settings
from openai import AsyncOpenAI
import numpy as np
from app.services.transcription_service import transcribe_media_file
from app.services.app_guide import build_app_guide

async def extract_text_from_file(file_path: str, filename: str) -> dict:
    """
    Returns {"text": str, "images": list[str]}
    Images are base64 data URIs: "data:image/png;base64,..."
    """
    ext = filename.split('.')[-1].lower()
    text = ""
    images = []
    try:
        if ext in ['jpg', 'jpeg', 'png', 'webp', 'gif']:
            with open(file_path, "rb") as image_file:
                base64_data = base64.b64encode(image_file.read()).decode('utf-8')
                mime_type = "image/jpeg" if ext in ['jpg', 'jpeg'] else f"image/{ext}"
                images.append(f"data:{mime_type};base64,{base64_data}")
        elif ext == 'pdf':
            with open(file_path, 'rb') as f:
                reader = pypdf.PdfReader(f)
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"

            # Image-based / scanned PDF (no real text layer): render each page
            # to a PNG so the multimodal model can read it via the images path.
            if len(text.strip()) < 20:
                try:
                    import fitz  # PyMuPDF
                    doc = fitz.open(file_path)
                    for page in doc[:10]:  # cap to first 10 pages
                        pix = page.get_pixmap(dpi=150)
                        b64 = base64.b64encode(pix.tobytes("png")).decode('utf-8')
                        images.append(f"data:image/png;base64,{b64}")
                    doc.close()
                except Exception as pdf_img_err:
                    print(f"PDF image-fallback failed for {filename}: {pdf_img_err}")
        elif ext in ['doc', 'docx']:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext == 'txt':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        elif ext in ['csv', 'xls', 'xlsx']:
            if ext == 'csv':
                df = pd.read_csv(file_path)
                text = df.to_string()
            else:
                # Read all sheets for text
                xls = pd.ExcelFile(file_path)
                sheet_texts = []
                for sheet_name in xls.sheet_names:
                    df = pd.read_excel(file_path, sheet_name=sheet_name)
                    sheet_texts.append(f"[Sheet: {sheet_name}]\n{df.to_string()}")
                text = "\n\n".join(sheet_texts)
                
                # Extract embedded images from Excel using openpyxl
                try:
                    from openpyxl import load_workbook
                    from io import BytesIO
                    wb = load_workbook(file_path)
                    for ws in wb.worksheets:
                        for img in ws._images:
                            try:
                                img_data = None
                                if hasattr(img, '_data'):
                                    img_data = img._data()
                                elif hasattr(img, 'ref'):
                                    img_data = img.ref.getvalue() if hasattr(img.ref, 'getvalue') else img.ref.read()
                                
                                if img_data:
                                    b64 = base64.b64encode(img_data).decode('utf-8')
                                    images.append(f"data:image/png;base64,{b64}")
                            except Exception as img_err:
                                print(f"Skipping image extraction: {img_err}")
                    wb.close()
                except Exception as excel_img_err:
                    print(f"Excel image extraction warning: {excel_img_err}")
        elif ext in ['mp3', 'wav', 'm4a', 'aac', 'ogg', 'flac', 'mp4', 'mov', 'avi', 'mkv', 'webm']:
            print(f"--- Transcribing Media Knowledge: {filename} ---")
            text = await transcribe_media_file(file_path)
            if not text:
                text = f"[Media File: {filename} - No speech detected or readable]"
    except Exception as e:
        print(f"Extraction Error for {filename}: {e}")
    return {"text": text, "images": images}

def chunk_text(text: str, chunk_size=800, overlap=150) -> list:
    chunks = []
    if not text: return chunks
    words = text.split()
    for i in range(0, len(words), chunk_size - overlap):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
    return chunks

async def process_knowledge_base(project_id: str, file_id: str, local_path: str, filename: str):
    """
    Background Task to index the knowledge base.
    """
    result = await extract_text_from_file(local_path, filename)
    text = result["text"]
    chunks = chunk_text(text)
    
    col = get_collection("KnowledgeBase")
    chunk_docs = []
    for c in chunks:
        chunk_docs.append({
            "project_id": project_id,
            "file_id": file_id,
            "filename": filename,
            "content": c,
            "created_at": datetime.utcnow()
        })
    
    if chunk_docs:
        await col.insert_many(chunk_docs)
    
    # Cleanup local file
    if os.path.exists(local_path):
        os.remove(local_path)

# Question-words and other stopwords that would otherwise match nearly every
# chunk and crowd out the genuinely informative keywords.
_KB_STOPWORDS = {
    "what", "when", "where", "which", "whose", "whom", "this", "that", "these",
    "those", "there", "their", "they", "them", "then", "than", "with", "from",
    "have", "does", "about", "tell", "give", "show", "list", "please", "could",
    "would", "should", "explain", "describe", "according", "uploaded", "file",
    "files", "document", "documents", "know", "want", "need", "help", "your",
    "into", "will", "shall", "been", "being", "were", "very", "much", "many",
    "some", "more", "most", "also", "just", "like", "make", "made", "based",
}


def _kb_keywords(query: str) -> list:
    """Informative search terms from a user question: word-tokenized (so trailing
    punctuation like "say?" doesn't poison the regex), stopword-filtered, escaped."""
    words = re.findall(r"[a-z0-9]+", (query or "").lower())
    return [re.escape(w) for w in words if len(w) > 3 and w not in _KB_STOPWORDS]


async def get_relevant_context(project_id: str, query: str, limit=5) -> str:
    """
    Simple keyword-based context retrieval (Mocking RAG for modularity).
    In a full production version, this would use vector embeddings.
    """
    col = get_collection("KnowledgeBase")
    keywords = _kb_keywords(query)

    chunks = []
    if keywords:
        regex = {"$regex": "|".join(keywords), "$options": "i"}
        # Content matches first — they are the genuinely relevant chunks.
        chunks = await col.find(
            {"project_id": project_id, "content": regex}
        ).limit(limit).to_list(limit)

        # Top up remaining slots from filename matches (e.g. "what does the
        # sales report say?" → Sales_Report.pdf) WITHOUT letting a filename hit
        # flood out content matches: filename matches every chunk of that file.
        if len(chunks) < limit:
            seen = [c["_id"] for c in chunks]
            extra = limit - len(chunks)
            chunks += await col.find(
                {"project_id": project_id, "filename": regex, "_id": {"$nin": seen}}
            ).limit(extra).to_list(extra)

    fell_back = False
    if not chunks:
        # No keyword hit (e.g. "summarize the uploaded document") — fall back to
        # the project's leading chunks instead of an empty context, so questions
        # the uploaded files DO cover aren't wrongly refused.
        chunks = await col.find({"project_id": project_id}).limit(limit).to_list(limit)
        fell_back = bool(chunks)

    context = "\n\n---\n\n".join([c["content"] for c in chunks])
    if fell_back:
        # Keep the "nothing matched" refusal signal the empty context used to
        # provide: the model should answer from these only if they truly cover
        # the question (they do for generic asks like "summarize this file").
        context = (
            "[Note: no snippet matched the question's keywords — these are the "
            "project's opening snippets. If they don't cover the question, say "
            "the knowledge base doesn't contain it.]\n\n" + context
        )
    return context

async def generate_ai_response(instructions: str, context: str, user_message: str, history: list, images: list = None, role: str = None, ctx=None):
    """Generate a Support Engine chat reply.

    `ctx` (an assistant UserContext) links this chat to the LMS: when provided,
    the model gets the assistant module's read-only, RBAC-scoped data tools and
    can answer live LMS questions (my sessions, my scores, my attendance, ...).
    Data scope is enforced server-side by each tool — never by the model.
    """
    try:
        client = AsyncOpenAI(api_key=settings.OPENAI_API_KEY)

        has_context = bool(context and context.strip())
        knowledge_block = context if has_context else "[No knowledge base content matched this question.]"
        app_guide = build_app_guide(role or (ctx.role if ctx else None))

        tools_schema = []
        if ctx is not None:
            # Lazy import: assistant modules import gpt_service (extractor,
            # attachment_store), so a module-level import here would be a cycle.
            from app.assistant.tools import registry
            registry.register_all()
            tools_schema = registry.openai_schema_for_role(ctx.role)

        if tools_schema:
            lms_block = (
                "You are connected to the Sparsh LMS through read-only data tools. "
                "Whenever the user asks about their (or, if their role permits, the "
                "organisation's) LIVE LMS records — profile, batches, courses/"
                "quarters, sessions, attendance, quiz/assessment results, learning "
                "progress, tasks, study recommendations, platform stats, their "
                "notifications, media-library files, Support Engine project access, "
                "session templates (staff), or recent activity (superadmin) — CALL "
                "the matching tool and answer from its result. Data scope is "
                "enforced server-side; never invent records, and if a tool returns "
                "no data, say so plainly."
            )
        else:
            lms_block = (
                "(No live data connection in this context — for questions about "
                "personal records, direct the user to the relevant Sparsh module.)"
            )

        system_prompt = f"""### SYSTEM INSTRUCTION
{instructions}

### KNOWLEDGE CONTEXT (uploaded files & data — ALWAYS answerable)
The following snippets are extracted from this project's uploaded knowledge base \
(documents, spreadsheets, media transcripts) AND any files or data the user has \
uploaded directly into this chat (those appear as "[Session Context - <filename>]" \
blocks, and uploaded images may accompany the user's message). They are your source \
of truth for questions about this content. Questions about ANY uploaded file or \
data — whatever its topic — are IN scope: answer them fully from these snippets.

---
{knowledge_block}
---

### SPARSH APPLICATION GUIDE (platform modules — ALWAYS answerable)
You may ALSO answer questions about the Sparsh platform itself and its modules \
(Dashboard, Companies, Batches, Quarters, Sessions & Calendar, Session Templates, \
Assessments/Quizzes, Training Roadmap, Reports, User & Team Management, Media \
Library, Support Engine, Notifications, Settings) — what they do, where to find \
them, and step-by-step how-to — using only the guide below. It describes how the \
application works; it is NOT live data, so never use it to state a user's actual \
records. Treat it as a second source of truth, used ONLY for questions about the \
platform and its modules.

---
{app_guide}
---

### LIVE LMS DATA (the user's records — answerable via tools)
{lms_block}

### RESPONSE PROTOCOL (STRICT GROUNDING — READ CAREFULLY)
These rules take precedence over the SYSTEM INSTRUCTION above: the instruction may \
set tone, persona, or domain emphasis, but it can NEVER widen what you are allowed \
to answer beyond the sources below.
1. **Answer ONLY from the sources above** — the KNOWLEDGE CONTEXT (this project's \
uploaded files and anything uploaded in this chat), the SPARSH APPLICATION GUIDE \
(the platform and its modules), or LIVE LMS DATA fetched through the provided \
tools. Every fact in your reply must be supported by one of them. Do NOT use your \
own general/training knowledge to answer, even if you are confident you know the \
answer.
2. **Uploaded content is always in scope**: if the user asks about a file or data \
they uploaded (in the project knowledge base or in this chat), answer it from the \
snippets — regardless of the topic of that file.
3. **LMS questions are in scope**: anything about the user's records or activity \
in this LMS (sessions, attendance, quiz scores, progress, batches, tasks, \
notifications, media-library files, Support Engine access) should be answered \
from the data tools when available — never from memory or assumption.
4. **Out-of-scope questions**: If NO source covers the question (e.g. general \
trivia, definitions, world knowledge, programming/tech concepts unrelated to \
Sparsh such as "what is ORM" as a database concept, or topics not in the uploaded \
files), do NOT answer it. Reply briefly that the question is outside this project's \
knowledge base and what Sparsh can help with, and invite a relevant question. Never \
substitute general knowledge.
5. **No fabrication**: Do not guess, infer beyond the sources, or fill gaps with \
assumptions. If only part of the answer is present, answer that part and say the \
rest isn't available. Never invent features, menu items, or records.
6. **Detail**: When the answer IS supported, be comprehensive, structured, and \
descriptive — but stay within what the sources support.
7. **Tone**: Maintain a professional, executive, and coaching-oriented tone.
8. **First messages & greetings**: these rules apply from the VERY FIRST message \
of a chat — if the user opens directly with an in-scope question (module how-to, \
uploaded content, or their LMS records), answer it immediately and fully; never \
ask them to rephrase or greet first. If the user only greets you, reply warmly in \
one line and offer what you can help with (this project's knowledge, the Sparsh \
platform, and their LMS records) — a greeting is never an out-of-scope question.
"""
        messages = [{"role": "system", "content": system_prompt}]
        # Add basic context from history
        for h in history[-5:]: # Last 5 messages for history
            messages.append({"role": h["role"], "content": h["content"]})
            
        user_content = [{"type": "text", "text": user_message}]
        
        if images:
            for img_url in images:
                user_content.append({
                    "type": "image_url",
                    # "auto" (not "low"): these are mostly scanned PDF pages and
                    # spreadsheet charts whose text is illegible when downscaled.
                    "image_url": {"url": img_url, "detail": "auto"}
                })
                
        messages.append({"role": "user", "content": user_content})

        # Tool-calling loop: lets the model fetch live LMS data (when ctx is
        # provided) before answering. Without tools this runs exactly once.
        for _ in range(4):
            kwargs = {"model": "gpt-4o", "messages": messages, "temperature": 0.7}
            if tools_schema:
                kwargs["tools"] = tools_schema
                kwargs["tool_choice"] = "auto"
            response = await client.chat.completions.create(**kwargs)
            msg = response.choices[0].message
            tool_calls = getattr(msg, "tool_calls", None)
            if not tool_calls:
                return msg.content

            from app.assistant.tools import registry
            messages.append({
                "role": "assistant",
                "content": msg.content,
                "tool_calls": [
                    {"id": c.id, "type": "function",
                     "function": {"name": c.function.name, "arguments": c.function.arguments}}
                    for c in tool_calls
                ],
            })
            for c in tool_calls:
                try:
                    args = json.loads(c.function.arguments or "{}")
                except json.JSONDecodeError:
                    args = {}
                spec = registry.get_tool(c.function.name)
                if spec is None:
                    payload = {"success": False, "error": "Unknown tool"}
                else:
                    # execute_tool re-checks the role and applies the server-side
                    # data scope from ctx — the model never widens access.
                    result = await registry.execute_tool(spec, ctx, args)
                    payload = result.for_llm()
                messages.append({
                    "role": "tool",
                    "tool_call_id": c.id,
                    "content": json.dumps(payload, default=str),
                })

        return ("I couldn't finish fetching the data needed to answer that. "
                "Please try asking again or narrowing the question.")
    except Exception as e:
        print(f"GPT Generation Error: {e}")
        return f"I'm sorry, I encountered an error: {str(e)}"

async def grade_descriptive_answer(question: str, user_answer: str, keywords: str, checker_instructions: str, max_marks: float):
    try:
        client = AsyncOpenAI(api_key=settings.OPENAI_API_KEY)
        
        prompt = f"""You are an Expert Academic Evaluator. Your task is to grade a student's answer based on the following criteria:

### QUESTION
{question}

### STUDENT'S ANSWER
{user_answer}

### EXPECTED CORE KEYWORDS / CONCEPTS
{keywords}

### SPECIAL INSTRUCTIONS FOR CHECKER
{checker_instructions}

### GRADING PROTOCOL
1. Assign a numeric score from 0 to {max_marks}.
2. Compare the student's answer against the keywords and question context.
3. Be fair but strict according to the special instructions.
4. Provide constructive feedback, explaining the REASON for the marks given and offering SUGGESTIONS for improvement.
5. If the answer is blank or irrelevant, assign 0.
6. Return ONLY a valid JSON object.

### OUTPUT FORMAT (JSON ONLY)
{{
  "score": <number>,
  "feedback": "<string>"
}}
"""
        response = await client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": prompt}],
            response_format={ "type": "json_object" },
            temperature=0.3
        )
        return json.loads(response.choices[0].message.content)
    except Exception as e:
        print(f"AI Grading Error: {e}")
        return {"score": 0, "feedback": f"AI Grading Engine Error: {str(e)}"}
