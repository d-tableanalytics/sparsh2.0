"""P1-143 verification harness -- the configuration snapshot.

"D-Table provides active Employee/Attendance/Payroll configuration and policy versions used
for UAT." Seeds one company's configuration across every area, builds the snapshot and checks
each section is there with its versions / effective dates, that the three renderings carry
it, and that building it WRITES NOTHING.

House convention: self-contained, no pytest, fake collections, ASCII output, exit 1 on fail.
Run:  python -m app.services.hrms.tests.test_config_snapshot   (from backend/)
"""
from __future__ import annotations

import asyncio
import io
import sys

from app.services.hrms.tests.test_phase2_employee import FakeCollection

results: list[bool] = []
CO = "C1"


def check(label: str, condition: bool, detail=None) -> bool:
    results.append(bool(condition))
    print(f"  {'PASS' if condition else 'FAIL'}  {label}"
          + ("" if condition or detail is None else f"  -- {detail}"))
    return bool(condition)


def section(title: str) -> None:
    print(f"\n-- {title} --")


WRITES: list = []


class Watched(FakeCollection):
    """A FakeCollection that records any write, so read-only can be asserted."""
    def __init__(self, name, docs=None):
        super().__init__(docs)
        self.name = name

    async def insert_one(self, doc):
        WRITES.append((self.name, "insert_one"))
        return await super().insert_one(doc)

    async def insert_many(self, docs):
        WRITES.append((self.name, "insert_many"))
        return await super().insert_many(docs)

    async def update_one(self, *a, **k):
        WRITES.append((self.name, "update_one"))
        return await super().update_one(*a, **k)

    async def update_many(self, *a, **k):
        WRITES.append((self.name, "update_many"))
        return await super().update_many(*a, **k)

    async def delete_one(self, *a, **k):
        WRITES.append((self.name, "delete_one"))
        return await super().delete_one(*a, **k)


async def main() -> int:
    import app.db.mongodb as mongo
    from app.models import hrms as M

    seed = {
        M.COLL_SETTINGS: [{
            "company_id": CO,
            "shift_policy": {"shift_start": "09:30", "shift_end": "18:30"},
            "shift_policy_version": 2,
            "shift_policy_versions": [
                {"version": 1, "effective_from": "2026-01-01", "by_name": "Hana HR",
                 "reason": "go-live", "old": {"shift_start": "09:00"},
                 "new": {"shift_start": "09:15"}},
                {"version": 2, "effective_from": "2026-03-01", "by_name": "Hana HR",
                 "reason": "later start", "old": {"shift_start": "09:15"},
                 "new": {"shift_start": "09:30"}}],
            "employee_master_lists": {"grades": ["M1", "E2"],
                                      "salary_templates": ["Staff 2026"]},
        }],
        M.COLL_SHIFTS: [{
            "company_id": CO, "code": "GEN", "name": "General", "active": True,
            "versions": [
                {"effective_from": "2026-01-01", "start": "09:30", "end": "18:30",
                 "break_minutes": 60, "min_hours": 8},
                {"effective_from": "2099-01-01", "start": "10:00", "end": "19:00",
                 "break_minutes": 60, "min_hours": 8.5}]}],
        M.COLL_HOLIDAYS: [{"company_id": CO, "holiday_date": "2026-03-25",
                           "holiday_name": "Holi (UAT)", "calendar": "Maharashtra",
                           "locations": ["Pune"]}],
        M.COLL_POLICY_DEFINITIONS: [{"company_id": CO, "code": "ATT_DAY",
                                     "module": "attendance", "name": "Day rules"}],
        M.COLL_POLICY_VERSIONS: [{"company_id": CO, "policy_code": "ATT_DAY",
                                  "version_no": 3, "status": "Active",
                                  "effective_from": "2026-02-01",
                                  "rules": {"half_day_fraction": 0.5}}],
        M.COLL_POLICY_ASSIGNMENTS: [{"company_id": CO, "policy_code": "ATT_DAY",
                                     "version_no": 3, "scope": {"location": "Pune"}}],
        M.COLL_LEAVE_TYPES: [{"company_id": CO, "code": "CL", "name": "Casual Leave",
                              "active": True, "annual_quota": 12}],
        M.COLL_SALARY_COMPONENTS: [{
            "company_id": CO, "code": "UAT_CONV", "name": "Conveyance",
            "component_type": "Earning", "calculation_method": "Formula",
            "formula": "MIN(2000, BASIC * 0.12)", "version": 2,
            "effective_from": "2025-04-01", "lifecycle": "Active", "active": True,
            "versions": [{"version": 1, "status": "Superseded",
                          "definition": {"calculation_method": "Formula",
                                         "formula": "BASIC * 0.12",
                                         "effective_from": "2025-01-01"}}]}],
        M.COLL_FIELD_DEFINITIONS: [{
            "company_id": CO, "module": "employee", "code": "uat_mgr_only",
            "label": "Manager only", "data_type": "Text", "status": "Active",
            "applicability": {"grade": ["M1"]}, "effective_from": "2026-01-01"}],
        M.COLL_PAYSLIP_TEMPLATE_VERSIONS: [{"company_id": CO, "version": 1,
                                            "company_name": "Sparsh Magic"}],
    }
    from bson import ObjectId
    store = {name: Watched(name, [{"_id": ObjectId(), **d} for d in docs])
             for name, docs in seed.items()}
    getter = lambda name: store.setdefault(name, Watched(name))  # noqa: E731
    mongo.get_collection = getter

    import app.services.hrms_config_snapshot_service as SNAP
    # The readers it reuses each import get_collection at module level; point them all at
    # the fakes (the employee master pulls in the shift master for its "shifts" list).
    import app.services.hrms_attendance_service       # noqa: F401
    import app.services.hrms_bank_format_service      # noqa: F401
    import app.services.hrms_config_change_service    # noqa: F401
    import app.services.hrms_employee_master_service  # noqa: F401
    import app.services.hrms_holiday_service          # noqa: F401
    import app.services.hrms_payroll_service          # noqa: F401
    import app.services.hrms_payslip_service          # noqa: F401
    import app.services.hrms_policy_version_service   # noqa: F401
    import app.services.hrms_shift_service            # noqa: F401
    for name, mod in list(sys.modules.items()):
        if name.startswith("app.services") and hasattr(mod, "get_collection"):
            mod.get_collection = getter

    section("1. Every area is in the snapshot")
    snap = await SNAP.build_snapshot(CO)
    by = {s["key"]: s for s in snap["sections"]}
    expected = ["attendance_rules", "attendance_rule_versions", "shifts", "holidays",
                "policy_versions", "leave_types", "payroll_policy", "salary_components",
                "salary_component_versions", "master_lists", "field_definitions",
                "bank_formats", "payslip_template", "payslip_template_versions",
                "governance"]
    check("all 15 sections present, in order", [s["key"] for s in snap["sections"]] == expected,
          [s["key"] for s in snap["sections"]])
    check("every section has columns and a title",
          all(s["columns"] and s["title"] for s in snap["sections"]))
    check("building the snapshot wrote nothing (read-only)", not WRITES, WRITES)

    section("2. Versions and effective dates travel with each area")
    rules = {r["setting"]: r["value"] for r in by["attendance_rules"]["rows"]}
    check("attendance rules: the version in force and the stored values",
          rules.get("(version in force)") == 2 and rules.get("shift_start") == "09:30", rules)
    hist = by["attendance_rule_versions"]["rows"]
    check("attendance rule history: both versions, oldest first, old -> new",
          [h["version"] for h in hist] == [1, 2]
          and "09:15 -> 09:30" in str(hist[1]["changes"]), hist)
    shifts = by["shifts"]["rows"]
    check("shifts: every version, the current one marked In force, the future Upcoming",
          [(s["effective_from"], s["state"]) for s in shifts]
          == [("2026-01-01", "In force"), ("2099-01-01", "Upcoming")], shifts)
    check("shifts: minimum hours per shift carried", shifts[0]["min_hours"] == 8)
    check("holidays: date, name and locations",
          by["holidays"]["rows"][0]["holiday_name"] == "Holi (UAT)"
          and by["holidays"]["rows"][0]["locations"] == "Pune", by["holidays"]["rows"])
    pv = by["policy_versions"]["rows"]
    check("policy engine: version, status, effective date, scope and rules",
          pv and pv[0]["version_no"] == 3 and pv[0]["status"] == "Active"
          and pv[0]["effective_from"] == "2026-02-01" and "Pune" in str(pv[0]["scope"])
          and "half_day_fraction" in str(pv[0]["rules"]), pv)
    check("leave types: code and their rules",
          by["leave_types"]["rows"][0]["code"] == "CL"
          and "annual_quota" in str(by["leave_types"]["rows"][0]["rules"]))
    comp = by["salary_components"]["rows"][0]
    check("salary component: the definition in force, with its version and dates",
          comp["code"] == "UAT_CONV" and comp["version"] == 2
          and comp["formula"] == "MIN(2000, BASIC * 0.12)"
          and comp["effective_from"] == "2025-04-01" and comp["proration"], comp)
    old = by["salary_component_versions"]["rows"]
    check("salary component history: v1 kept as Superseded with its own formula",
          old and old[0]["version"] == 1 and old[0]["status"] == "Superseded"
          and old[0]["formula"] == "BASIC * 0.12", old)
    check("payroll policy section has settings",
          len(by["payroll_policy"]["rows"]) > 0)
    lists = {r["list"]: r["values"] for r in by["master_lists"]["rows"]}
    check("master lists incl. salary structure templates",
          lists.get("grades") == "M1; E2" and lists.get("salary_templates") == "Staff 2026",
          lists)
    fd = by["field_definitions"]["rows"][0]
    check("custom field: status, applicability and effective date",
          fd["code"] == "uat_mgr_only" and fd["status"] == "Active"
          and "M1" in str(fd["applicability"]) and fd["effective_from"] == "2026-01-01", fd)
    check("bank formats: DEFAULT present even when none configured",
          any(f["name"] == "DEFAULT" for f in by["bank_formats"]["rows"]))
    check("payslip template versions listed",
          by["payslip_template_versions"]["rows"][0]["version"] == 1)
    gov = {g["type"]: g for g in by["governance"]["rows"]}
    check("governance: every known type, with needs-approval and cannot-switch-off",
          gov and all(g["governed"] in (True, False) for g in gov.values())
          and any(g["always"] for g in gov.values()), gov)

    section("3. The three renderings")
    body, media, name = await SNAP.export(CO, "xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(body))
    check("xlsx: an About sheet plus one sheet per section",
          wb.sheetnames == ["About"] + expected, wb.sheetnames)
    check("xlsx: a section sheet carries its rows under a header",
          wb["salary_components"]["A2"].value == "UAT_CONV"
          and wb["salary_components"]["A1"].value == "Code")
    check("xlsx: media type and file name",
          media.endswith("spreadsheetml.sheet") and name.endswith(".xlsx"))
    body, media, name = await SNAP.export(CO, "csv")
    text = body.decode("utf-8-sig")
    check("csv: every section titled in turn",
          all(f"## {s['title']}" in text for s in snap["sections"]) and name.endswith(".csv"))
    js = await SNAP.export(CO, "json")
    check("json: the snapshot itself", isinstance(js, dict) and len(js["sections"]) == 15)
    check("still nothing written after three exports", not WRITES, WRITES)

    section("4. Wiring")
    from app.routes import hrms as routes
    paths = {getattr(r, "path", "") for r in routes.router.routes}
    check("GET /config/snapshot is registered",
          any(p.endswith("/config/snapshot") for p in paths))

    passed = sum(results)
    print(f"\n=== {passed}/{len(results)} checks passed ===")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
