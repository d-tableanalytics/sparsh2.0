"""HRMS > configuration snapshot (P1-143).

"D-Table provides active Employee/Attendance/Payroll configuration and policy versions used
for UAT." The configuration audit (P1-128) says who changed what; this says what IS in force,
in one document: one section per area, every row carrying its version and effective dates
where the area has them.

-- Read-only by construction --------------------------------------------------------------
Every section is built from the module's own readers (list_shifts, list_salary_components,
list_field_definitions, ...), so the snapshot shows exactly what the screens and the
calculators see -- the same defaults filled in, the same DEFAULT bank format synthesised.
The one exception is leave types: `list_leave_types` SEEDS the defaults on first read, and
an export must never write, so that section reads the collection directly.

-- Shape ----------------------------------------------------------------------------------
`build_snapshot` returns {"sections": [{key, title, columns, rows}]}, rows as dicts. JSON
returns that as is; XLSX writes one sheet per section behind an "About" sheet; CSV stacks
the sections with a title line between them, since a CSV has no sheets.
"""
import csv
import io
import json
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import HTTPException

from app.db.mongodb import get_collection
from app.models.hrms import COLL_LEAVE_TYPES, COLL_SETTINGS

# A nested value (a rule set, a slab table) in one cell. Excel's cell limit is 32,767.
_CELL_LIMIT = 32000


def _compact(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (list, tuple)) and all(
            isinstance(v, (str, int, float)) for v in value):
        return "; ".join(str(v) for v in value)
    if isinstance(value, dict) and not value:
        return None
    text = json.dumps(value, default=str, sort_keys=True, ensure_ascii=False)
    return text if len(text) <= _CELL_LIMIT else text[:_CELL_LIMIT] + " ...(truncated)"


def _section(key: str, title: str, columns: list, rows: list) -> dict:
    return {"key": key, "title": title, "columns": columns,
            "rows": [{k: _compact(r.get(k)) for k, _ in columns} for r in rows]}


_KV = [("setting", "Setting"), ("value", "Value")]


def _kv_rows(values: dict, skip=("_id", "company_id")) -> list:
    return [{"setting": k, "value": v} for k, v in sorted((values or {}).items())
            if k not in skip]


def _today() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


# ─────────────────────────────────────────────────────────────
# Sections
# ─────────────────────────────────────────────────────────────
async def _attendance(company_id: str) -> list:
    from app.services.hrms_attendance_service import get_shift_policy
    policy = await get_shift_policy(company_id)
    settings = await get_collection(COLL_SETTINGS).find_one(
        {"company_id": str(company_id)},
        {"shift_policy_version": 1, "shift_policy_versions": 1}) or {}
    rules = _kv_rows(policy)
    rules.insert(0, {"setting": "(version in force)",
                     "value": settings.get("shift_policy_version") or "default (never changed)"})
    history = [{**v, "changes": {k: f"{(v.get('old') or {}).get(k)} -> {n}"
                                 for k, n in (v.get("new") or {}).items()}}
               for v in sorted(settings.get("shift_policy_versions") or [],
                               key=lambda v: int(v.get("version") or 0))]
    return [
        _section("attendance_rules", "Attendance rules (company policy in force)", _KV, rules),
        _section("attendance_rule_versions", "Attendance rule versions",
                 [("version", "Version"), ("effective_from", "Effective From"),
                  ("by_name", "Changed By"), ("approved_by", "Approved By"),
                  ("via", "Via"), ("reason", "Reason"), ("changes", "Changes (old -> new)")],
                 history),
    ]


async def _shifts(company_id: str) -> dict:
    from app.services.hrms_shift_service import list_shifts
    today = _today()
    rows = []
    for s in await list_shifts(company_id):
        current = (s.get("current") or {}).get("effective_from")
        for v in sorted(s.get("versions") or [], key=lambda v: v.get("effective_from") or ""):
            state = ("In force" if v.get("effective_from") == current
                     else "Upcoming" if (v.get("effective_from") or "") > today else "Past")
            rows.append({**v, "code": s.get("code"), "name": s.get("name"),
                         "active": s.get("active", True), "state": state,
                         "applicability": s.get("applicability")})
    return _section("shifts", "Shifts (every version)",
                    [("code", "Shift"), ("name", "Name"), ("active", "Active"),
                     ("effective_from", "Effective From"), ("state", "State"),
                     ("start", "Start"), ("end", "End"), ("break_minutes", "Break (min)"),
                     ("min_hours", "Min Hours"), ("daily_grace_minutes", "Late Grace (min)"),
                     ("early_going_grace_minutes", "Early Grace (min)"),
                     ("half_day_threshold_minutes", "Half Day Below (min)"),
                     ("absent_threshold_minutes", "Absent Below (min)"),
                     ("monthly_buffer_minutes", "Monthly Buffer (min)"),
                     ("flexible", "Flexible"), ("weekly_off_pattern", "Weekly Off"),
                     ("applicability", "Applies To"), ("reason", "Reason")],
                    rows)


async def _holidays(company_id: str) -> dict:
    from app.services.hrms_holiday_service import list_holidays
    rows = (await list_holidays(company_id)).get("holidays") or []
    return _section("holidays", "Holiday calendar",
                    [("holiday_date", "Date"), ("holiday_name", "Holiday"),
                     ("calendar", "Calendar"), ("locations", "Locations")], rows)


async def _policy_versions(company_id: str) -> dict:
    """The versioned attendance / payroll / day policies of the policy engine (§28)."""
    from app.services.hrms_policy_version_service import list_definitions, list_versions
    rows = []
    for d in await list_definitions(company_id):
        for v in await list_versions(company_id, d["code"]):
            rows.append({**v, "code": d.get("code"), "module": d.get("module"),
                         "category": d.get("category"), "policy_name": d.get("name"),
                         "scope": [a.get("scope") or {k: a.get(k) for k in a
                                                      if k not in ("id", "company_id")}
                                   for a in v.get("assignments") or []]})
    return _section("policy_versions", "Attendance / payroll policy versions",
                    [("code", "Policy"), ("module", "Module"), ("category", "Category"),
                     ("policy_name", "Name"), ("version_no", "Version"), ("status", "Status"),
                     ("effective_from", "Effective From"), ("effective_to", "Effective To"),
                     ("approved_by_name", "Approved By"), ("scope", "Assigned To"),
                     ("change_reason", "Reason"), ("rules", "Rules")], rows)


async def _leave_types(company_id: str) -> dict:
    # Read directly: list_leave_types seeds defaults on first read, and an export never writes.
    docs = await get_collection(COLL_LEAVE_TYPES).find(
        {"company_id": str(company_id)}).sort("code", 1).to_list(100)
    skip = {"_id", "company_id", "code", "name", "active", "policy_confirmed",
            "created_at", "updated_at"}
    rows = [{**d, "rules": {k: v for k, v in d.items() if k not in skip}} for d in docs]
    return _section("leave_types", "Leave types",
                    [("code", "Code"), ("name", "Name"), ("active", "Active"),
                     ("policy_confirmed", "Policy Confirmed"), ("rules", "Rules")], rows)


async def _payroll(company_id: str) -> list:
    from app.services.hrms_payroll_service import get_payroll_policy, list_salary_components
    policy = await get_payroll_policy(company_id, on_date=_today())
    components = await list_salary_components(None, company_id)
    history = []
    for c in components:
        for v in c.get("versions") or []:
            d = v.get("definition") or {}
            history.append({**d, "code": c.get("code"), "version": v.get("version"),
                            "status": v.get("status")})
    calc = [("calculation_method", "Method"), ("calculation_base", "Base"),
            ("percentage", "Percentage"), ("formula", "Formula"), ("slabs", "Slabs")]
    return [
        _section("payroll_policy", "Payroll policy in force (incl. approval controls)",
                 _KV, _kv_rows(policy)),
        _section("salary_components", "Salary components (definition in force)",
                 [("code", "Code"), ("name", "Name"), ("component_type", "Type"), *calc,
                  ("proration", "Proration"), ("rounding", "Rounding"),
                  ("version", "Version"), ("effective_from", "Effective From"),
                  ("effective_to", "Effective To"), ("lifecycle", "Lifecycle"),
                  ("active", "Active"), ("statutory", "Statutory"),
                  ("payslip_section", "Payslip Section"), ("sequence", "Sequence"),
                  ("frequency", "Frequency"), ("pending_request_no", "Pending Change")],
                 components),
        _section("salary_component_versions", "Salary component versions (superseded)",
                 [("code", "Code"), ("version", "Version"), ("status", "Status"), *calc,
                  ("effective_from", "Effective From"), ("effective_to", "Effective To")],
                 history),
    ]


async def _employee_config(company_id: str) -> list:
    from app.services.hrms_employee_master_service import (
        get_master_lists, list_field_definitions,
    )
    lists = await get_master_lists(company_id)
    fields = await list_field_definitions(company_id)
    return [
        _section("master_lists", "Employee master lists (incl. salary structure templates)",
                 [("list", "List"), ("values", "Values")],
                 [{"list": k, "values": v} for k, v in sorted(lists.items())]),
        _section("field_definitions", "Custom field definitions",
                 [("module", "Module"), ("code", "Code"), ("label", "Label"),
                  ("data_type", "Type"), ("status", "Status"), ("section", "Section"),
                  ("mandatory", "Mandatory"), ("sensitive", "Sensitive"),
                  ("applicability", "Applies To"), ("effective_from", "Effective From"),
                  ("effective_to", "Effective To"), ("visible_to_self", "Visible To Self"),
                  ("visible_to_manager", "Visible To Manager"),
                  ("visible_to_payroll", "Visible To Payroll")], fields),
    ]


async def _outputs(company_id: str) -> list:
    from app.services.hrms_bank_format_service import list_formats
    from app.services.hrms_payslip_service import get_payslip_template, list_template_versions
    formats = [{**f, "column_labels": [c.get("label") or c.get("source")
                                       for c in f.get("columns") or []]}
               for f in await list_formats(company_id)]
    template = await get_payslip_template(company_id)
    versions = await list_template_versions(company_id)
    return [
        _section("bank_formats", "Bank file formats",
                 [("name", "Format"), ("bank", "Bank"), ("version", "Version"),
                  ("built_in", "Built In"), ("delimiter", "Delimiter"),
                  ("fixed_width", "Fixed Width"), ("date_format", "Date Format"),
                  ("amount_format", "Amount Format"), ("include_totals_row", "Totals Row"),
                  ("column_labels", "Columns")], formats),
        _section("payslip_template", "Payslip template in force", _KV, _kv_rows(template)),
        _section("payslip_template_versions", "Payslip template versions",
                 [("version", "Version"), ("created_at", "Saved At"),
                  ("updated_by", "Saved By (user id)"), ("company_name", "Company Name")],
                 versions),
    ]


async def _governance(company_id: str) -> dict:
    from app.services.hrms_config_change_service import get_governed_types
    g = await get_governed_types(company_id)
    governed, always = set(g["governed_config_types"]), set(g["always_governed"])
    rows = [{"type": t, "governed": t in governed, "always": t in always}
            for t in g["known"]]
    return _section("governance", "Configuration types needing approval (FRD §27)",
                    [("type", "Configuration Type"), ("governed", "Needs Approval"),
                     ("always", "Cannot Be Switched Off")], rows)


async def build_snapshot(company_id: str) -> dict:
    """Every area's active configuration for one company. Reads only."""
    if not company_id:
        raise HTTPException(status_code=422, detail="Choose a company.")
    sections = []
    sections += await _attendance(company_id)
    sections.append(await _shifts(company_id))
    sections.append(await _holidays(company_id))
    sections.append(await _policy_versions(company_id))
    sections.append(await _leave_types(company_id))
    sections += await _payroll(company_id)
    sections += await _employee_config(company_id)
    sections += await _outputs(company_id)
    sections.append(await _governance(company_id))
    return {"company_id": str(company_id),
            "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            "sections": sections}


# ─────────────────────────────────────────────────────────────
# Rendering
# ─────────────────────────────────────────────────────────────
def render_csv(snapshot: dict) -> bytes:
    """The sections one after another: a title line, the header, the rows, a blank line."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([f"HRMS configuration snapshot - company {snapshot['company_id']} - "
                f"generated {snapshot['generated_at']}"])
    for s in snapshot["sections"]:
        w.writerow([])
        w.writerow([f"## {s['title']}"])
        w.writerow([label for _, label in s["columns"]])
        for r in s["rows"]:
            w.writerow(["" if r.get(k) is None else r.get(k) for k, _ in s["columns"]])
        if not s["rows"]:
            w.writerow(["(none configured)"])
    return buf.getvalue().encode("utf-8-sig")


def render_xlsx(snapshot: dict) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="Excel export is unavailable on this server. Please export CSV.")
    wb = Workbook()
    about = wb.active
    about.title = "About"
    about.append(["HRMS configuration snapshot"])
    about["A1"].font = Font(bold=True, size=14)
    about.append(["Company", snapshot["company_id"]])
    about.append(["Generated", snapshot["generated_at"]])
    about.append([])
    about.append(["Sheet", "Section", "Rows"])
    for cell in about[5]:
        cell.font = Font(bold=True)
    names = set()
    for s in snapshot["sections"]:
        # Sheet names: 31 characters, unique, no []:*?/\ -- the key already satisfies the rest.
        name = s["key"][:31]
        while name in names:
            name = name[:29] + "_" + str(len(names))
        names.add(name)
        about.append([name, s["title"], len(s["rows"])])
        ws = wb.create_sheet(name)
        ws.append([label for _, label in s["columns"]])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F46E5")
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for r in s["rows"]:
            ws.append([r.get(k) for k, _ in s["columns"]])
        if not s["rows"]:
            ws.append(["(none configured)"])
        for i, (_k, label) in enumerate(s["columns"], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(
                60, max(12, len(label) + 2))
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(company_id: str, fmt: str) -> str:
    return f"hrms_configuration_snapshot_{_today()}.{fmt}"


async def export(company_id: str, fmt: Optional[str] = "xlsx"):
    """(body, media type, file name) -- or the snapshot dict itself for fmt="json"."""
    snapshot = await build_snapshot(company_id)
    if fmt == "json":
        return snapshot
    if fmt == "csv":
        return render_csv(snapshot), "text/csv; charset=utf-8", file_name(company_id, "csv")
    return (render_xlsx(snapshot),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_name(company_id, "xlsx"))
