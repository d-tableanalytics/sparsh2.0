"""HRMS > Attendance > capture/import layer and timing rules (client requirement).

"Biometric login/logout/live-location practice, office timing 9:30-6:30, grace/buffer
references, outdoor duty and flexible-timing approval" -> "Attendance capture/import layer,
OD, regularisation, timing rules and audit."

This module holds the pieces the rest of attendance did not have:

  * EFFECTIVE HOURS -- the hours a day is judged against: an approved flexible timing for
    that date if there is one, otherwise the company's office timing (shift policy). Every
    writer of a day (HR's mark, self check-in, regularisation, biometric import) goes
    through `compute_day`, so a day is never judged against the wrong hours.
  * FLEXIBLE TIMING -- an employee asks for different hours over a date range (e.g. 10:30 to
    19:30); the reporting manager / HR approves (the same decision rule as OD and
    regularisation: never your own). Approving re-judges any day already recorded in range.
  * BIOMETRIC IMPORT -- the machine's export (CSV or Excel). First punch of a day = check-in,
    last = check-out. Never overwrites a locked day, an approved leave / OD / holiday /
    weekly off, or a regularised day; everything skipped is reported with its reason.
  * MONTHLY BUFFER -- each person's late minutes in a month are tolerated up to the
    buffer; the Late Coming view shows what is within it and what is beyond.

-- Where a day's rules come from (FRD §23, §26) -------------------------------------------
Four layers, most specific first. Each supplies only the keys it actually sets, so a
company that configures one thing does not have to restate the rest:

    1. an approved FLEXIBLE TIMING for that employee on that date (start/end only)
    2. the ATT_DAY POLICY VERSION in force on that date for that employee
    3. the SHIFT version in force on that date (shift master, already versioned)
    4. the company shift policy (hrms_settings) -- where every company started

Layer 2 is new here. Nothing above or below it changed, and a company with no ATT_DAY
policy resolves exactly as it did before this phase: that is not a hope, it is what
`resolve_rules(..., default=...)` does when there is no version, and the suite pins it.
"""
import csv
import io
import re
from datetime import date, datetime, timedelta, timezone
from typing import Optional

from fastapi import HTTPException

from app.db.mongodb import get_collection
from app.models.hrms import (
    COLL_ATTENDANCE, COLL_EMPLOYEE_PROFILES, COLL_PUNCH_SEGMENTS, compute_daily_status,
)
from app.services.hrms_audit_service import audit
from app.services.hrms_id_service import next_business_id

from app.models.hrms import COLL_FLEXI_TIMING as COLL_FLEXI
ENTITY_FLEXI = "flexi_timing"
ENTITY_ATTENDANCE = "attendance"
AUDIT_FLEXI_REQUESTED = "flexible timing requested"
AUDIT_FLEXI_ACTIONED = "flexible timing decided"
AUDIT_BIOMETRIC_IMPORT = "biometric attendance imported"

FLEXI_PENDING, FLEXI_APPROVED, FLEXI_REJECTED, FLEXI_CANCELLED = (
    "Pending", "Approved", "Rejected", "Cancelled")
# Statuses a biometric file must never overwrite: somebody decided these.
PROTECTED_STATUSES = {"On Leave", "On OD", "Holiday", "Weekly Off"}
PROTECTED_SOURCES = {"regularization", "od"}
MAX_IMPORT_ROWS = 20000
MAX_FLEXI_DAYS = 92

_HHMM = re.compile(r"^([01]?\d|2[0-3]):([0-5]\d)$")


def _hhmm(value, label: str) -> str:
    text = str(value or "").strip()
    m = _HHMM.match(text)
    if not m:
        raise HTTPException(status_code=422, detail=f"{label} must be a time like 09:30.")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def _iso(value, label: str) -> str:
    try:
        return date.fromisoformat(str(value).strip()[:10]).isoformat()
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail=f"{label} must be a date (YYYY-MM-DD).")


# ─────────────────────────────────────────────────────────────
# Effective hours for one person on one day
# ─────────────────────────────────────────────────────────────
async def effective_hours(company_id: str, employee_code: str, work_date: str,
                          policy: dict = None) -> dict:
    """The rules a day is judged against, and where each came from.

    See the module docstring for the four layers. The returned `policy` is the flattened
    result every caller already reads; `rules` keeps the shift provenance and
    `day_policy_version` the ATT_DAY one, so a day can record WHICH versions judged it
    (§25 Audit, §30 reproducibility).
    """
    if policy is None:
        from app.services.hrms_attendance_service import get_shift_policy
        policy = await get_shift_policy(company_id)
    flexi = await get_collection(COLL_FLEXI).find_one({
        "company_id": str(company_id), "employee_code": employee_code,
        "status": FLEXI_APPROVED, "from_date": {"$lte": work_date},
        "to_date": {"$gte": work_date}})
    # FRD §5 / §23: the employee's own shift (or the policy that applies to them), in the
    # version in force on this date -- falling back to the company timing.
    from app.services.hrms_shift_service import profile_for, resolve_rules
    profile = await profile_for(company_id, employee_code)
    rules = await resolve_rules(company_id, profile, work_date, policy=policy)
    merged = {**policy, **{k: v for k, v in rules.items() if v is not None}}

    # FRD §23/§26 -- the ATT_DAY policy version for this employee on this date, laid over
    # the shift. Only keys it actually SETS are applied: a policy that configures the
    # half-day combination must not blank out the grace the shift already gives.
    from app.services import hrms_policy_version_service as PV
    from app.models.hrms import POLICY_ATT_DAY
    day_version = await PV.resolve(company_id, POLICY_ATT_DAY, profile or {}, work_date)
    if day_version:
        merged.update({k: v for k, v in (day_version.get("rules") or {}).items()
                       if v is not None})

    if flexi:
        return {"start": flexi["shift_start"], "end": flexi["shift_end"],
                "flexi_no": flexi.get("flexi_no"), "policy": merged, "rules": rules,
                "day_policy_version": day_version}
    return {"start": rules["shift_start"], "end": rules["shift_end"], "flexi_no": None,
            "policy": merged, "rules": rules, "day_policy_version": day_version}


async def compute_day(company_id: str, employee_code: str, work_date: str,
                      actual_in: Optional[str], actual_out: Optional[str],
                      policy: dict = None, *, approved_exception: Optional[str] = None) -> dict:
    """compute_daily_status against the day's EFFECTIVE hours. Adds `flexi_no` when a
    flexible timing decided them, and the shift / rule version that judged the day."""
    hours = await effective_hours(company_id, employee_code, work_date, policy)
    p = hours["policy"]
    out = compute_daily_status(
        scheduled_in=hours["start"], scheduled_out=hours["end"],
        actual_in=actual_in, actual_out=actual_out,
        daily_grace_minutes=p["daily_grace_minutes"],
        half_day_threshold_minutes=p["half_day_threshold_minutes"],
        early_going_grace_minutes=p.get("early_going_grace_minutes") or 0,
        break_minutes=p.get("break_minutes") or 0,
        min_hours=p.get("min_hours"),
        absent_threshold_minutes=p.get("absent_threshold_minutes"),
        # FRD §23 -- configured where a policy sets them, today's behaviour where not.
        half_day_fraction=p.get("half_day_fraction", 0.5),
        half_day_late_minutes=p.get("half_day_late_minutes"),
        half_day_early_minutes=p.get("half_day_early_minutes"),
        absent_on_no_out_punch=bool(p.get("absent_on_no_out_punch")))
    # §23 Half Day "approved exception treatment": a late/early half day is waived when an
    # approved request of a type the policy names covers the day. Off unless configured.
    waived_types = p.get("half_day_exception_types") or []
    if approved_exception and approved_exception in waived_types \
            and out.get("half_day_basis") in ("late", "early"):
        out.update({"status": "Present", "payable_day": 1.0, "lop_fraction": 0.0,
                    "half_day_reason": f"waived: approved {approved_exception}"})
    # §5 / §23 Late Coming "late threshold": minutes beyond grace are always measured, but a
    # day is a LATE MARK only past the threshold. Unset, any late minute is one, as before.
    threshold = int(p.get("late_threshold_minutes") or 0)
    out["late_mark"] = int(out.get("late_minutes") or 0) > threshold
    out["flexi_no"] = hours["flexi_no"]
    out["scheduled_in"], out["scheduled_out"] = hours["start"], hours["end"]
    rules = hours.get("rules") or {}
    # A FLEXIBLE shift has core hours, not a start bell: the day is judged by the hours
    # worked (half-day / minimum-hours rules above still apply), never as late or early.
    if rules.get("flexible") and not hours["flexi_no"]:
        out["late_mark"] = False
        out["late_minutes"] = 0
        out["early_minutes"] = 0
        out["grace_applied_minutes"] = 0
        out["flexible"] = True
    out["shift_code"] = rules.get("shift_code")
    out["rule_source"] = "flexible timing" if hours["flexi_no"] else rules.get("rule_source")
    out["rule_version"] = rules.get("shift_version")
    # §30: "Every payroll result records sufficient rule/component version references to
    # reproduce the result later." A day is the payroll's input, so it carries the same
    # debt -- which ATT_DAY version judged it, if any did.
    day_version = hours.get("day_policy_version") or {}
    out["day_policy_version"] = day_version.get("version_no")
    out["day_policy_code"] = day_version.get("policy_code")
    return out


DERIVED_DAY_FIELDS = ("scheduled_in", "scheduled_out", "early_minutes", "grace_applied_minutes",
                      "payable_minutes", "payable_hours", "scheduled_minutes", "payable_day",
                      "lop_fraction", "missing_out_punch", "shift_code", "rule_source",
                      "rule_version", "half_day_reason", "day_policy_code",
                      "day_policy_version", "late_mark")


def derived_fields(computed: dict) -> dict:
    """The FRD §6 'Calculated' / 'Payroll' fields every writer of a day stores alongside
    status / worked / late minutes."""
    return {k: computed.get(k) for k in DERIVED_DAY_FIELDS if k in computed}


# ─────────────────────────────────────────────────────────────
# Flexible timing
# ─────────────────────────────────────────────────────────────
async def request_flexi(actor: dict, company_id: str, payload: dict) -> dict:
    from app.services.hrms_attendance_service import (
        _assert_self_or_privileged, _get_profile, _reporting_manager_id)
    code = str(payload.get("employee_code") or "").strip()
    if not code:
        raise HTTPException(status_code=422, detail="Choose the employee.")
    profile = await _get_profile(company_id, code)
    await _assert_self_or_privileged(actor, company_id, code)
    start_d = _iso(payload.get("from_date"), "From date")
    end_d = _iso(payload.get("to_date") or payload.get("from_date"), "To date")
    if end_d < start_d:
        raise HTTPException(status_code=422, detail="The end date is before the start date.")
    if (date.fromisoformat(end_d) - date.fromisoformat(start_d)).days + 1 > MAX_FLEXI_DAYS:
        raise HTTPException(status_code=422,
                            detail=f"Ask for at most {MAX_FLEXI_DAYS} days at a time.")
    s_in, s_out = _hhmm(payload.get("shift_start"), "Start time"), _hhmm(payload.get("shift_end"), "End time")
    if s_out <= s_in:
        raise HTTPException(status_code=422, detail="The end time must be after the start time.")
    reason = str(payload.get("reason") or "").strip()[:1000]
    if not reason:
        raise HTTPException(status_code=422, detail="Say why the different hours are needed.")
    clash = await get_collection(COLL_FLEXI).find_one({
        "company_id": str(company_id), "employee_code": code,
        "status": {"$in": [FLEXI_PENDING, FLEXI_APPROVED]},
        "from_date": {"$lte": end_d}, "to_date": {"$gte": start_d}}, {"flexi_no": 1})
    if clash:
        raise HTTPException(status_code=409,
                            detail=f"{clash['flexi_no']} already covers some of these dates.")

    now = datetime.now(timezone.utc)
    flexi_no = await next_business_id("flexi", str(company_id), now.year)
    doc = {
        "flexi_no": flexi_no, "company_id": str(company_id), "employee_code": code,
        "employee_name": profile.get("display_name") or profile.get("full_name"),
        "reporting_manager_id": await _reporting_manager_id(profile),
        "from_date": start_d, "to_date": end_d, "shift_start": s_in, "shift_end": s_out,
        "reason": reason, "status": FLEXI_PENDING,
        "requested_by": str(actor.get("_id") or ""), "action_by": None, "action_at": None,
        "remarks": None, "created_at": now, "updated_at": now,
    }
    await get_collection(COLL_FLEXI).insert_one(dict(doc))
    await audit(actor, AUDIT_FLEXI_REQUESTED, ENTITY_FLEXI, flexi_no,
                f"{code} {start_d}..{end_d} {s_in}-{s_out}", company_id)
    doc.pop("_id", None)
    return doc


async def list_flexi(actor: dict, company_id: str, *, status: str = None, limit: int = 200) -> list:
    from app.services.hrms_attendance_service import _scope_query
    query = {"company_id": str(company_id)}
    if status:
        query["status"] = status
    query = await _scope_query(actor, company_id, query)
    rows = await get_collection(COLL_FLEXI).find(query).sort("created_at", -1).to_list(limit)
    for r in rows:
        r.pop("_id", None)
    return rows


async def act_on_flexi(actor: dict, company_id: str, flexi_no: str, payload: dict) -> dict:
    from app.services.hrms_attendance_service import _assert_may_decide
    coll = get_collection(COLL_FLEXI)
    doc = await coll.find_one({"company_id": str(company_id), "flexi_no": flexi_no})
    if not doc:
        raise HTTPException(status_code=404, detail=f"Flexible timing '{flexi_no}' not found.")
    if doc["status"] != FLEXI_PENDING:
        raise HTTPException(status_code=409, detail=f"{flexi_no} is already \"{doc['status']}\".")
    await _assert_may_decide(actor, company_id, doc["employee_code"], "flexible timing request")
    approved = bool(payload.get("approved"))
    remarks = str(payload.get("remarks") or "").strip()[:1000] or None
    if not approved and not remarks:
        raise HTTPException(status_code=422, detail="Say why it is not approved.")
    now = datetime.now(timezone.utc)
    new_status = FLEXI_APPROVED if approved else FLEXI_REJECTED
    await coll.update_one({"company_id": str(company_id), "flexi_no": flexi_no},
                          {"$set": {"status": new_status, "action_by": str(actor.get("_id") or ""),
                                    "action_at": now, "remarks": remarks, "updated_at": now}})
    rejudged = 0
    if approved:
        rejudged = await _rejudge_range(company_id, doc["employee_code"],
                                        doc["from_date"], doc["to_date"])
    await audit(actor, AUDIT_FLEXI_ACTIONED, ENTITY_FLEXI, flexi_no,
                f"{new_status}; {rejudged} recorded day(s) re-judged", company_id)
    out = await coll.find_one({"company_id": str(company_id), "flexi_no": flexi_no})
    out.pop("_id", None)
    out["rejudged_days"] = rejudged
    return out


async def _rejudge_range(company_id: str, code: str, start_d: str, end_d: str) -> int:
    """Re-judge days already recorded in a range against their (now flexible) hours."""
    rows = await get_collection(COLL_ATTENDANCE).find({
        "company_id": str(company_id), "employee_code": code,
        "work_date": {"$gte": start_d, "$lte": end_d}}).to_list(200)
    n = 0
    for r in rows:
        if r.get("locked") or r.get("status") in PROTECTED_STATUSES or not r.get("actual_in"):
            continue
        c = await compute_day(company_id, code, r["work_date"], r.get("actual_in"), r.get("actual_out"))
        await get_collection(COLL_ATTENDANCE).update_one(
            {"company_id": str(company_id), "employee_code": code, "work_date": r["work_date"]},
            {"$set": {"status": c["status"], "late_minutes": c["late_minutes"],
                      "worked_minutes": c["worked_minutes"], "flexi_no": c["flexi_no"],
                      **derived_fields(c), "updated_at": datetime.now(timezone.utc)}})
        n += 1
    return n


# ─────────────────────────────────────────────────────────────
# Biometric import
# ─────────────────────────────────────────────────────────────
_CODE_COLS = ("employee code", "emp code", "employee id", "emp id", "empcode", "employee_code",
              "code", "card no", "biometric id", "enroll no", "user id")
_DATE_COLS = ("date", "work date", "attendance date", "punch date", "work_date")
_TIME_COLS = ("time", "punch time", "punch", "log time")
_DATETIME_COLS = ("datetime", "date time", "punch datetime", "timestamp", "log date")
_IN_COLS = ("in", "in time", "check in", "first in", "in_time", "actual_in")
_OUT_COLS = ("out", "out time", "check out", "last out", "out_time", "actual_out")


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " "))


def _pick(headers: list, names) -> Optional[int]:
    wanted = {_norm(n) for n in names}
    for i, h in enumerate(headers):
        if _norm(h) in wanted:
            return i
    return None


def _parse_date(v) -> Optional[str]:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    text = str(v or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text[:11].strip(), fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _parse_time(v) -> Optional[str]:
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if hasattr(v, "hour") and hasattr(v, "minute"):
        return f"{v.hour:02d}:{v.minute:02d}"
    text = str(v or "").strip()
    m = re.search(r"(\d{1,2}):(\d{2})(?::\d{2})?\s*([AaPp][Mm])?", text)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    ap = (m.group(3) or "").lower()
    if ap == "pm" and h < 12:
        h += 12
    if ap == "am" and h == 12:
        h = 0
    return f"{h:02d}:{mi:02d}" if h < 24 and mi < 60 else None


def _parse_datetime(v):
    if isinstance(v, datetime):
        return v.date().isoformat(), v.strftime("%H:%M")
    text = str(v or "").strip()
    return _parse_date(text[:10]) or _parse_date(text.split(" ")[0]), _parse_time(text)


def read_rows(filename: str, raw: bytes) -> list:
    """The file as a list of rows (first row = headers). CSV or Excel."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    if name.endswith(".csv") or name.endswith(".txt") or not name:
        text = raw.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    raise HTTPException(status_code=415, detail="Upload the biometric export as CSV or Excel (.xlsx).")


def punches_from_rows(rows: list) -> tuple:
    """{(code, date): [times]} plus a list of unreadable rows. Pure."""
    if not rows:
        raise HTTPException(status_code=422, detail="The file is empty.")
    headers = rows[0]
    ci = _pick(headers, _CODE_COLS)
    if ci is None:
        raise HTTPException(
            status_code=422,
            detail="No employee code column found. Name it e.g. \"Employee Code\".")
    di, ti = _pick(headers, _DATE_COLS), _pick(headers, _TIME_COLS)
    dti = _pick(headers, _DATETIME_COLS)
    ii, oi = _pick(headers, _IN_COLS), _pick(headers, _OUT_COLS)
    if dti is None and di is None:
        raise HTTPException(status_code=422,
                            detail="No date column found. Name it \"Date\" (or \"Date Time\").")
    if dti is None and ti is None and ii is None and oi is None:
        raise HTTPException(status_code=422,
                            detail="No time column found. Use \"Time\", or \"In Time\" / \"Out Time\".")
    days, bad = {}, []
    for n, row in enumerate(rows[1:MAX_IMPORT_ROWS + 1], start=2):
        if not row or all(str(c or "").strip() == "" for c in row):
            continue
        cell = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
        code = str(cell(ci) or "").strip()
        if isinstance(cell(ci), float) and float(cell(ci)).is_integer():
            code = str(int(cell(ci)))
        if dti is not None:
            d, t = _parse_datetime(cell(dti))
            times = [t] if t else []
        else:
            d = _parse_date(cell(di))
            times = [x for x in (_parse_time(cell(ti)) if ti is not None else None,
                                 _parse_time(cell(ii)) if ii is not None else None,
                                 _parse_time(cell(oi)) if oi is not None else None) if x]
        if not code or not d or not times:
            bad.append({"row": n, "reason": "missing employee code, date or time"})
            continue
        days.setdefault((code, d), []).extend(times)
    return days, bad


def mapped_values_from_rows(rows: list, mapping: dict) -> dict:
    """{(code, date): {field_code: value}} for the columns attendance fields are mapped to
    (FRD §21 "source mapping"). The last non-empty value of a day wins. Pure."""
    if not rows or not mapping:
        return {}
    headers = rows[0]
    ci, di = _pick(headers, _CODE_COLS), _pick(headers, _DATE_COLS)
    dti = _pick(headers, _DATETIME_COLS)
    columns = {field: _pick(headers, [column]) for field, column in mapping.items()}
    columns = {f: i for f, i in columns.items() if i is not None}
    out = {}
    if ci is None or (di is None and dti is None) or not columns:
        return out
    for row in rows[1:MAX_IMPORT_ROWS + 1]:
        cell = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
        code = str(cell(ci) or "").strip()
        if isinstance(cell(ci), float) and float(cell(ci)).is_integer():
            code = str(int(cell(ci)))
        d = _parse_datetime(cell(dti))[0] if dti is not None else _parse_date(cell(di))
        if not code or not d:
            continue
        for field, i in columns.items():
            value = cell(i)
            if value not in (None, ""):
                out.setdefault((code, d), {})[field] = value
    return out


async def import_biometric(actor: dict, company_id: str, filename: str, raw: bytes) -> dict:
    """Load a biometric export. First punch = in, last punch = out."""
    rows = read_rows(filename, raw)
    days, bad = punches_from_rows(rows)
    # §21 source mapping: attendance fields that name a column of this file.
    from app.services.hrms_employee_master_service import (
        list_field_definitions, validate_custom_values)
    mapping = {d["code"]: d["source_column"]
               for d in await list_field_definitions(company_id, "attendance")
               if d.get("status") == "Active" and d.get("source_column")}
    mapped = mapped_values_from_rows(rows, mapping)
    profiles = await get_collection(COLL_EMPLOYEE_PROFILES).find(
        {"company_id": str(company_id)}, {"employee_code": 1, "biometric_id": 1}).to_list(5000)
    by_code = {}
    for p in profiles:
        if p.get("employee_code"):
            by_code[str(p["employee_code"]).upper()] = p["employee_code"]
        if p.get("biometric_id"):
            by_code[str(p["biometric_id"]).upper()] = p["employee_code"]

    from app.services.hrms_attendance_service import get_shift_policy
    policy = await get_shift_policy(company_id)
    now = datetime.now(timezone.utc)
    report = {"days_created": 0, "days_updated": 0, "skipped": [], "unreadable_rows": bad,
              "unknown_codes": [], "field_errors": []}
    unknown = set()
    for (raw_code, d), times in sorted(days.items()):
        code = by_code.get(raw_code.upper())
        if not code:
            unknown.add(raw_code)
            continue
        times = sorted(set(times))
        actual_in = times[0]
        actual_out = times[-1] if len(times) > 1 else None
        key = {"company_id": str(company_id), "employee_code": code, "work_date": d}
        existing = await get_collection(COLL_ATTENDANCE).find_one(key) or {}
        if existing.get("locked"):
            report["skipped"].append({"employee_code": code, "work_date": d, "reason": "locked"})
            continue
        if existing.get("status") in PROTECTED_STATUSES or existing.get("source") in PROTECTED_SOURCES:
            report["skipped"].append({"employee_code": code, "work_date": d,
                                      "reason": f"already {existing.get('status') or existing.get('source')}"})
            continue
        if not existing and await _month_locked(company_id, d):
            report["skipped"].append({"employee_code": code, "work_date": d, "reason": "locked"})
            continue
        c = await compute_day(company_id, code, d, actual_in, actual_out, policy)
        custom = {}
        if mapped.get((raw_code, d)):
            try:
                custom = {"custom_fields": await validate_custom_values(
                    company_id, "attendance", mapped[(raw_code, d)],
                    current=existing.get("custom_fields"), enforce_mandatory=False)}
            except HTTPException as e:
                report["field_errors"].append({"employee_code": code, "work_date": d,
                                               "reason": str(e.detail)})
        await get_collection(COLL_ATTENDANCE).update_one(
            key,
            {"$set": {**key, **custom, "actual_in": actual_in, "actual_out": actual_out,
                      "status": c["status"], "worked_minutes": c["worked_minutes"],
                      "late_minutes": c["late_minutes"], "flexi_no": c["flexi_no"],
                      **derived_fields(c),
                      "source": "biometric", "source_reference": filename,
                      "locked": False, "updated_at": now,
                      "updated_by": str(actor.get("_id") or "")},
             "$setOnInsert": {"created_at": now, "original_in": actual_in,
                              "original_out": actual_out, "original_source": "biometric"}},
            upsert=True)
        await get_collection(COLL_PUNCH_SEGMENTS).insert_one({
            **key, "punch_in": actual_in, "punch_out": actual_out, "all_punches": times,
            "source": "biometric", "file": filename, "recorded_by": str(actor.get("_id") or ""),
            "recorded_at": now})
        report["days_updated" if existing else "days_created"] += 1
    report["unknown_codes"] = sorted(unknown)
    await audit(actor, AUDIT_BIOMETRIC_IMPORT, ENTITY_ATTENDANCE, filename or "biometric file",
                f"{report['days_created']} new, {report['days_updated']} updated, "
                f"{len(report['skipped'])} skipped, {len(unknown)} unknown code(s)", company_id)
    return report


async def _month_locked(company_id: str, work_date: str) -> bool:
    from app.models.hrms import COLL_ATTENDANCE_LOCKS
    lock = await get_collection(COLL_ATTENDANCE_LOCKS).find_one(
        {"company_id": str(company_id), "period": str(work_date)[:7]})
    return (lock or {}).get("status") == "Locked"


# ─────────────────────────────────────────────────────────────
# Monthly buffer
# ─────────────────────────────────────────────────────────────
async def late_escalation_thresholds(company_id: str, profiles: dict, rows: list) -> dict:
    """FRD §23 Late Coming "escalation threshold": {(employee_code, 'YYYY-MM'): late days}.

    An ATT_DAY rule, so it is resolved per person through the policy spine like every other
    late-coming rule, as at the month's last day (or today, for the month in progress). A
    company that has never set it pays one query and gets {} -- nothing is flagged.
    """
    from app.models.hrms import COLL_POLICY_VERSIONS, POLICY_ATT_DAY
    configured = await get_collection(COLL_POLICY_VERSIONS).find_one(
        {"company_id": str(company_id), "policy_code": POLICY_ATT_DAY,
         "rules.late_escalation_days": {"$gt": 0}}, {"_id": 1})
    if not configured:
        return {}
    import calendar
    from app.services import hrms_policy_version_service as PV
    today = date.today().isoformat()
    out = {}
    for code, month in {(r["employee_code"], str(r["work_date"])[:7]) for r in rows}:
        year, mon = int(month[:4]), int(month[5:7])
        on = min(f"{month}-{calendar.monthrange(year, mon)[1]:02d}", today)
        version = await PV.resolve(company_id, POLICY_ATT_DAY, profiles.get(code) or {}, on)
        days = ((version or {}).get("rules") or {}).get("late_escalation_days")
        if days:
            out[(code, month)] = int(days)
    return out


def monthly_buffer(rows: list, buffer_minutes: int, escalation: dict = None) -> list:
    """Per person per month: late days, late minutes, and how much is within / beyond the
    monthly buffer -- plus, where an escalation threshold applies (see
    late_escalation_thresholds), whether the late days have reached it. Pure."""
    buckets = {}
    for r in rows:
        key = (r["employee_code"], str(r["work_date"])[:7])
        b = buckets.setdefault(key, {"employee_code": key[0], "month": key[1], "late_days": 0,
                                     "late_minutes": 0, "employee_name": r.get("employee_name")})
        b["late_days"] += 1
        b["late_minutes"] += int(r.get("late_minutes") or 0)
    out = []
    for b in buckets.values():
        b["buffer_minutes"] = int(buffer_minutes or 0)
        b["within_buffer"] = min(b["late_minutes"], b["buffer_minutes"])
        b["beyond_buffer"] = max(0, b["late_minutes"] - b["buffer_minutes"])
        threshold = (escalation or {}).get((b["employee_code"], b["month"]))
        b["escalation_days"] = threshold
        b["escalated"] = bool(threshold) and b["late_days"] >= threshold
        out.append(b)
    return sorted(out, key=lambda x: (x["month"], -x["beyond_buffer"], x["employee_code"]), reverse=False)
