"""
TPMS ▸ per-client report export.

One client's complete TPMS record as a single .xlsx, structured for analysis rather than for
re-import. This is deliberately NOT tpms_backup_service: that one dumps raw collections keyed
by ObjectId so the workbook can be read back in, which is the right shape for a backup and the
wrong shape for someone who wants to pivot a client's year in Excel or Google Sheets. Here the
ids are resolved to names, the statuses are spelled out, and every sheet is a flat table with a
header row you can filter on.

WHAT IS IN IT
-------------
  Summary              the selected period's KPIs — the same figures the Client View shows
  Success Measures     per-activity targets, actuals and achievement, ALL periods
  Scheduled Activities every TPMS activity scheduled for the client, with completion + delay
  Activity Tracker     per-person occurrence rows, with Done / Missed / Pending derived
  Escalations          active and resolved, with level and ageing
  Action Items         owner, target date, delay, status
  Reschedule Requests  old → new date with the reason given
  Uploads              proof files attached against activities
  People               the client's roster, for grouping the sheets above by person

Only the Summary is period-scoped (it mirrors the on-screen dashboard). Every detail sheet
carries the client's FULL history plus a Period / Date column, so a month can be filtered in
the spreadsheet without having to re-export — which is what "complete data" has to mean if the
report is going to be analysed rather than just read.

READ-ONLY: nothing here writes, and no collection is created or altered.
"""
from __future__ import annotations

import io
import logging
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from app.db.mongodb import get_collection
from app.models.tpms import (
    COLL_ACTION_ITEMS, COLL_ACTIVITY_TRACKER, COLL_ESCALATIONS,
    COLL_RESCHEDULE_REQUESTS, COLL_SUCCESS_MEASURES, COLL_TASK_UPLOADS,
    TPMS_EVENT_KIND, period_display,
)
from app.controllers.auth_controller import CLIENT_RANK, client_rank, is_client_side_user
from app.services.tpms_dashboard_service import (
    STAFF_ROLES, _allowed_companies, _classify, _load_companies, _period_window,
    get_learner_dashboard,
)
from app.services.tpms_schedule_service import CAL_COLLECTIONS

logger = logging.getLogger(__name__)

SCHEDULE_COLLECTIONS = list(dict.fromkeys(CAL_COLLECTIONS + ["calendar_events"]))
MAX_ROWS = 50000


# ─────────────────────────────────────────────────────────────
# Cell helpers
# ─────────────────────────────────────────────────────────────
def _txt(v: Any) -> Any:
    """A value Excel can hold. Dates become real datetimes so they sort and filter as dates;
    everything unprintable becomes a string so openpyxl never raises mid-write."""
    if v is None:
        return ""
    if isinstance(v, (int, float, bool, datetime)):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (list, tuple, set)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={x}" for k, x in v.items())
    return str(v)


def _day(v: Any) -> str:
    """The YYYY-MM-DD part of whatever the document stored — ISO string or datetime."""
    if not v:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    return str(v)[:10]


def _days_between(later: Any, earlier: Any) -> Any:
    """Whole days `later - earlier`, or "" when either side is unusable. Used for delay and
    ageing columns, where a blank is meaningfully different from a 0."""
    a, b = _day(later), _day(earlier)
    if not a or not b:
        return ""
    try:
        return (date.fromisoformat(a) - date.fromisoformat(b)).days
    except ValueError:
        return ""


# ─────────────────────────────────────────────────────────────
# Data gathering
# ─────────────────────────────────────────────────────────────
async def _people(company_id: str) -> Dict[str, dict]:
    """Everyone who can appear on this client's rows, keyed by id. Staff are merged in because
    an activity's coach / scheduler is an internal user, and an unresolved id in a report
    column is exactly the thing this export exists to avoid."""
    out: Dict[str, dict] = {}
    for coll, side in (("learners", "Client"), ("staff", "Internal")):
        query = {"company_id": str(company_id)} if coll == "learners" else {}
        for u in await get_collection(coll).find(query).to_list(5000):
            name = (u.get("full_name")
                    or " ".join(filter(None, [u.get("first_name"), u.get("last_name")])).strip()
                    or u.get("email") or str(u["_id"]))
            out[str(u["_id"])] = {
                "name": name,
                "email": u.get("email") or "",
                "governance_role": u.get("governance_role") or "",
                "designation": u.get("designation") or "",
                "department": u.get("department") or "",
                "active": u.get("is_active", True),
                "side": side,
            }
    return out


async def _schedules(company_id: str) -> List[dict]:
    """TPMS activities live in the ERP calendar collections tagged `kind: tpms_activity`, not
    in a table of their own — so they are gathered by scan across all three."""
    rows: List[dict] = []
    for coll in SCHEDULE_COLLECTIONS:
        try:
            rows += await get_collection(coll).find({
                "kind": TPMS_EVENT_KIND, "company_id": str(company_id),
            }).to_list(MAX_ROWS)
        except Exception as exc:  # noqa: BLE001 — a missing collection must not fail the report
            logger.warning("TPMS client export: could not read %s — %s", coll, exc)
    rows.sort(key=lambda r: _day(r.get("start")))
    return rows


# ─────────────────────────────────────────────────────────────
# Sheet builders — each returns (title, columns, rows, note)
# ─────────────────────────────────────────────────────────────
def _sheet_success(measures: List[dict]) -> tuple:
    cols = ["Period", "Activity", "Scope", "HOD",
            "Implementation Target %", "Implementation Actual %",
            "Score Target %", "Score Actual %", "Achievement %", "Last Updated"]
    rows = [[
        period_display(str(m.get("period") or "")) or m.get("period"),
        m.get("activity"), m.get("scope"), m.get("hod_name"),
        m.get("impl_target"), m.get("impl_actual"),
        m.get("score_target"), m.get("score_actual"), m.get("achievement"),
        m.get("updated_at"),
    ] for m in sorted(measures, key=lambda x: (str(x.get("period") or ""),
                                               str(x.get("activity") or "").lower()))]
    return ("Success Measures", cols, rows,
            "Achievement % = Score Actual ÷ Score Target × 100. A blank Achievement means the "
            "activity has no score yet — it is NOT a zero.")


def _sheet_schedules(schedules: List[dict], people: Dict[str, dict]) -> tuple:
    cols = ["Date", "Period", "Activity", "Title", "Status", "Assigned To",
            "Completed On", "Delay (days)", "Doer Marked Done", "Doer Marked On",
            "Reschedules", "Scheduled By", "Batch ID"]
    rows = []
    for s in schedules:
        day = _day(s.get("start"))
        completed = s.get("completed_at")
        # Only a LATE completion is a delay; finishing early is not negative delay, it is 0.
        delay = _days_between(completed, day) if completed else ""
        if isinstance(delay, int) and delay < 0:
            delay = 0
        members = ", ".join(people.get(str(m), {}).get("name", str(m))
                            for m in (s.get("assigned_member_ids") or []))
        rows.append([
            day, day[:7], s.get("activity"), s.get("title"),
            s.get("tpms_status") or s.get("status"), members,
            _day(completed), delay,
            "Yes" if s.get("learner_done") else "", _day(s.get("learner_done_at")),
            s.get("reschedule_count") or 0,
            people.get(str(s.get("user_id")), {}).get("name", ""),
            s.get("tpms_batch_id"),
        ])
    return ("Scheduled Activities", cols, rows,
            "One row per scheduled occurrence. Delay counts only completions later than the "
            "scheduled date; Cancelled occurrences are listed but are not counted as planned "
            "work anywhere in TPMS.")


def _sheet_tracker(tracker: List[dict], people: Dict[str, dict]) -> tuple:
    cols = ["Date", "Period", "Activity", "Person", "Designation", "Raw Status", "Outcome"]
    today = date.today().isoformat()
    rows = []
    for t in sorted(tracker, key=lambda x: (_day(x.get("date")), str(x.get("activity") or ""))):
        person = people.get(str(t.get("member_id")), {})
        day = _day(t.get("date"))
        status = str(t.get("status") or "")
        # Same derivation the HOD / Employee dashboards use — the tracker never stores
        # "Missed", so without this the column would simply be absent from the data.
        outcome = {"done": "Done", "missed": "Missed", "pending": "Pending"}[
            _classify(status, day, today)]
        rows.append([day, t.get("period"), t.get("activity"),
                     person.get("name", str(t.get("member_id") or "")),
                     person.get("designation", ""), status, outcome])
    return ("Activity Tracker", cols, rows,
            "Outcome is derived: Missed = status Lapsed OR the date passed without "
            "completion. Cancelled rows are excluded from TPMS scoring.")


def _sheet_escalations(escalations: List[dict]) -> tuple:
    cols = ["Activity", "Level", "Escalated To", "Escalation Date", "Target Date",
            "Status", "Days Overdue", "Last Reminder", "Recommended Action", "OM"]
    today = date.today().isoformat()
    rows = []
    for e in sorted(escalations, key=lambda x: _day(x.get("escalation_date")), reverse=True):
        resolved = str(e.get("status") or "") == "Resolved"
        overdue = "" if resolved else _days_between(today, e.get("target_date"))
        if isinstance(overdue, int) and overdue < 0:
            overdue = 0
        rows.append([e.get("activity"), e.get("level"), e.get("escalated_to"),
                     _day(e.get("escalation_date")), _day(e.get("target_date")),
                     e.get("status") or "Active", overdue,
                     _day(e.get("last_reminder")), e.get("recommended_action"), e.get("om")])
    return ("Escalations", cols, rows,
            "Days Overdue is measured against the TARGET date and floored at 0; it is left "
            "blank once an escalation is Resolved.")


def _sheet_actions(actions: List[dict]) -> tuple:
    cols = ["Action", "Activity", "Owner", "Owner Email", "Target Date",
            "Status", "Delay (days)", "Raised On"]
    rows = [[a.get("action"), a.get("activity"), a.get("owner_name"), a.get("owner_email"),
             _day(a.get("target_date")), a.get("status"), a.get("delay_days"),
             _day(a.get("created_at"))]
            for a in sorted(actions, key=lambda x: _day(x.get("target_date")))]
    return ("Action Items", cols, rows,
            "Action Closure % across TPMS = items with status 'Closed' ÷ all items × 100.")


def _sheet_reschedules(requests: List[dict]) -> tuple:
    cols = ["Activity", "Title", "From Date", "From Time", "To Date", "To Time",
            "Reason", "Requested By", "Requested On", "Status", "Decided By", "Decided On",
            "Note"]
    rows = [[r.get("activity"), r.get("title"), _day(r.get("old_date")), r.get("old_time"),
             _day(r.get("new_date")), r.get("new_time"), r.get("reason"),
             r.get("requested_by_name"), _day(r.get("requested_at")), r.get("status"),
             r.get("decided_by"), _day(r.get("decided_at")), r.get("note")]
            for r in sorted(requests, key=lambda x: _day(x.get("requested_at")), reverse=True)]
    return ("Reschedule Requests", cols, rows, "")


def _sheet_uploads(uploads: List[dict]) -> tuple:
    cols = ["Uploaded On", "Period", "Activity", "Scope", "File", "Size (KB)",
            "Uploaded By", "For Person"]
    rows = [[_day(u.get("uploaded_at")), u.get("period"), u.get("activity"), u.get("scope"),
             u.get("file_name"),
             round((u.get("size") or 0) / 1024, 1) if u.get("size") else "",
             u.get("uploaded_by_name"), u.get("member_name")]
            for u in sorted(uploads, key=lambda x: _day(x.get("uploaded_at")), reverse=True)]
    return ("Uploads", cols, rows,
            "Proof files attached against activities. The files themselves stay in storage — "
            "this sheet is the index.")


def _sheet_people(people: Dict[str, dict]) -> tuple:
    cols = ["Name", "Email", "Governance Role", "Designation", "Department", "Side", "Active"]
    rows = [[p["name"], p["email"], p["governance_role"], p["designation"],
             p["department"], p["side"], "Yes" if p["active"] else "No"]
            for p in sorted(people.values(), key=lambda x: (x["side"], x["name"].lower()))
            if p["side"] == "Client"]
    return ("People", cols, rows,
            "The client's roster. Governance Role drives who may assign to whom "
            "(MD > HR > HOD > Implementor).")


# ─────────────────────────────────────────────────────────────
# Authorization
# ─────────────────────────────────────────────────────────────
class ExportNotPermitted(Exception):
    """The caller may not export this client."""


async def assert_may_export(user: dict, company_id: str) -> None:
    """Gate the export to companies the caller is actually entitled to.

    Checked HERE rather than relying on the dashboard's scoping, because
    `_allowed_companies` returns an explicitly-passed company_id as-is — it treats the
    parameter as already-authorised, which is survivable for a single dashboard read and is
    not survivable for a bulk export of a client's entire record.

      client MD / clientadmin → their own company only
      admin/superadmin        → any company
      other internal          → the companies they own as OM

    A plain `clientuser` is refused outright: the workbook is the WHOLE company's record —
    every person's tracker rows, every escalation — which is the company admin's to take out,
    not an individual doer's. They keep their own dashboard either way.
    """
    company_id = str(company_id)
    if is_client_side_user(user):
        if str(user.get("company_id") or "") != company_id:
            raise ExportNotPermitted("You can only export your own company's report.")
        if client_rank(user) < CLIENT_RANK["MD"]:
            raise ExportNotPermitted(
                "Only your company's admin can download the company report.")
        return
    if (user.get("role") or "").lower() in STAFF_ROLES:
        return
    companies = await _load_companies()
    owned = _allowed_companies(user, companies, {})   # empty scope → the caller's true set
    if owned is not None and company_id not in owned:
        raise ExportNotPermitted("You can only export reports for your own clients.")


# ─────────────────────────────────────────────────────────────
# Assembly
# ─────────────────────────────────────────────────────────────
async def build_client_report(user: dict, company_id: str,
                              period: Optional[str] = None) -> dict:
    """Gather one client's whole TPMS record plus the selected period's headline KPIs.

    The KPIs come from get_learner_dashboard — the very function the Client View renders — so
    the Summary sheet and the screen the user clicked Download on cannot disagree.
    """
    company_id = str(company_id)
    await assert_may_export(user, company_id)
    window = _period_window(period)
    period_key = period or window[0][:7]

    company = await get_collection("companies").find_one({"_id": _oid(company_id)}) or {}
    people = await _people(company_id)

    dashboard: dict = {}
    try:
        dashboard = await get_learner_dashboard(user, {"period": period_key,
                                                       "company_id": company_id}) or {}
    except Exception as exc:  # noqa: BLE001 — the detail sheets are still worth exporting
        logger.warning("TPMS client export: dashboard KPIs unavailable — %s", exc)

    scoped = {"company_id": company_id}
    measures = await get_collection(COLL_SUCCESS_MEASURES).find(scoped).to_list(MAX_ROWS)
    tracker = await get_collection(COLL_ACTIVITY_TRACKER).find(scoped).to_list(MAX_ROWS)
    escalations = await get_collection(COLL_ESCALATIONS).find(scoped).to_list(MAX_ROWS)
    actions = await get_collection(COLL_ACTION_ITEMS).find(scoped).to_list(MAX_ROWS)
    reschedules = await get_collection(COLL_RESCHEDULE_REQUESTS).find(scoped).to_list(MAX_ROWS)
    uploads = await get_collection(COLL_TASK_UPLOADS).find(scoped).to_list(MAX_ROWS)
    schedules = await _schedules(company_id)

    sheets = [
        _sheet_success(measures),
        _sheet_schedules(schedules, people),
        _sheet_tracker(tracker, people),
        _sheet_escalations(escalations),
        _sheet_actions(actions),
        _sheet_reschedules(reschedules),
        _sheet_uploads(uploads),
        _sheet_people(people),
    ]

    return {
        "company_id": company_id,
        "company": company.get("name") or dashboard.get("company") or company_id,
        "om": dashboard.get("om") or "",
        "period": period_key,
        "period_label": period_display(period_key),
        "period_from": window[0],
        "period_to": window[1],
        "generated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "generated_by": user.get("full_name") or user.get("email") or "",
        "dashboard": dashboard,
        "sheets": sheets,
    }


def _oid(value):
    from bson import ObjectId
    from bson.errors import InvalidId
    try:
        return ObjectId(str(value))
    except (InvalidId, TypeError):
        return value


def export_client_workbook(report: dict) -> bytes:
    """The report as one .xlsx — a Summary sheet, then one flat table per dataset."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4F46E5")
    title_font = Font(bold=True, size=14)
    label_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    dash = report.get("dashboard") or {}
    op = dash.get("op_cards") or {}
    cards = dash.get("cards") or {}

    # ── Summary ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"TPMS Report — {report.get('company')}"])
    ws["A1"].font = title_font
    ws.append([])
    for label, value in (
        ("Client", report.get("company")),
        ("Operations Manager", report.get("om") or "—"),
        ("Reporting period", f"{report.get('period_label') or report.get('period')} "
                             f"({report.get('period_from')} → {report.get('period_to')})"),
        ("Generated at (UTC)", report.get("generated_at")),
        ("Generated by", report.get("generated_by")),
    ):
        ws.append([label, _txt(value)])
        ws.cell(row=ws.max_row, column=1).font = label_font

    ws.append([])
    ws.append([f"Key figures — {report.get('period_label') or report.get('period')}"])
    ws.cell(row=ws.max_row, column=1).font = label_font
    ws.append(["Metric", "Value", "How it is calculated"])
    for col in (1, 2, 3):
        c = ws.cell(row=ws.max_row, column=col)
        c.font, c.fill = head_font, head_fill
    for label, value, how in (
        ("Planned activities", op.get("planned"), "Scheduled occurrences in the period, excluding Cancelled"),
        ("Completed activities", op.get("completed"), "Occurrences with status Completed"),
        ("Completion %", op.get("completion"), "Completed ÷ Planned × 100"),
        ("Average delay (days)", op.get("avg_delay"), "Σ days late ÷ number of late completions"),
        ("Health band", dash.get("status"), "≥95 STRONG · ≥85 GOOD · ≥70 WATCH · else AT-RISK"),
        ("Success measures tracked", cards.get("total"), "Activities on the scorecard this period"),
        ("Met", cards.get("met"), "Achievement % ≥ 100"),
        ("Partial", cards.get("partial"), "Achievement % between 50 and 99"),
        ("Not Met", cards.get("not_met"), "Achievement % below 50, or no score recorded"),
        ("Average achievement %", cards.get("avg_score"), "Σ Achievement % ÷ activities that have one"),
    ):
        ws.append([label, _txt(value if value is not None else "—"), how])

    ws.append([])
    ws.append(["Contents"])
    ws.cell(row=ws.max_row, column=1).font = label_font
    ws.append(["Sheet", "Rows", "Notes"])
    for col in (1, 2, 3):
        c = ws.cell(row=ws.max_row, column=col)
        c.font, c.fill = head_font, head_fill
    for title, _cols, rows, note in report.get("sheets") or []:
        ws.append([title, len(rows), note])

    ws.append([])
    ws.append(["The figures above cover the selected period only. Every detail sheet carries "
               "this client's FULL history with a Period / Date column, so any month can be "
               "filtered in the sheet without re-exporting."])
    ws.cell(row=ws.max_row, column=1).alignment = wrap
    for col, width in zip("ABC", (30, 46, 70)):
        ws.column_dimensions[col].width = width

    # ── One sheet per dataset ────────────────────────────────
    for title, cols, rows, note in report.get("sheets") or []:
        sheet = wb.create_sheet(title[:31])
        sheet.append(cols)
        for idx in range(1, len(cols) + 1):
            c = sheet.cell(row=1, column=idx)
            c.font, c.fill, c.alignment = head_font, head_fill, wrap
        for row in rows:
            sheet.append([_txt(v) for v in row])
        # Filter + frozen header: the point of the export is slicing it in the spreadsheet.
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        for idx, name in enumerate(cols, start=1):
            longest = max([len(str(name))] + [len(str(r[idx - 1])) for r in rows[:200]] or [0])
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 12), 46)

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
